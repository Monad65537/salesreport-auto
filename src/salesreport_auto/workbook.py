"""Restrained openpyxl rendering and full logical workbook read-back checks."""

import io
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.data_source import AxDataSource, StrRef
from openpyxl.chart.series import SeriesLabel
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_to_tuple
from openpyxl.worksheet.table import Table, TableStyleInfo

from .errors import ReportError
from .report_model import (
    CHART_INTEGER_LIMIT,
    CellValue,
    ChartPlan,
    SheetPlan,
    WorkbookPlan,
    sheet_names,
)

EXCEL_MAX_ROWS = 1_048_576
EXCEL_MAX_TEXT = 32_767
_NAVY = "243B53"
_BLUE = "E8EFF5"
_INK = "243746"
_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def _check_value(value: CellValue) -> None:
    if isinstance(value, str):
        try:
            length = len(value.encode("utf-16-le")) // 2
        except UnicodeError as exc:
            raise ReportError("Report contains text unsupported by Excel/XML") from exc
        if length > EXCEL_MAX_TEXT:
            raise ReportError("Report text exceeds Excel's 32,767 UTF-16-unit cell limit")
        if ILLEGAL_CHARACTERS_RE.search(value) or "\ufffe" in value or "\uffff" in value:
            raise ReportError("Report contains text unsupported by Excel/XML")
    elif value is not None and type(value) not in (int, bool):
        raise ReportError("Report cells must contain text, exact integers, booleans, or blanks")
    elif type(value) is int and abs(value) > CHART_INTEGER_LIMIT:
        raise ReportError("Large integer report cells must be stored as text")


def _put(sheet, row: int, column: int, value: CellValue) -> None:
    try:
        _check_value(value)
    except ReportError as exc:
        raise ReportError(f"{sheet.title}, cell {get_column_letter(column)}{row}: {exc}") from exc
    if value is None:
        return
    cell = sheet.cell(row, column, value)
    cell.font = Font(name="Calibri", size=11, color=_INK)
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    if isinstance(value, str):
        cell.data_type = "s"  # Including formula-like IDs and authoritative money.
        cell.number_format = "@"
    elif type(value) is int:
        cell.number_format = "#,##0;[Red](#,##0);0"


def _render_chart(sheet, plan: ChartPlan) -> None:
    chart = LineChart() if plan.kind == "monthly" else BarChart()
    if plan.kind in ("branches", "products"):
        chart.type = "bar" if plan.kind == "products" else "col"
        chart.grouping = "clustered"
    if plan.kind == "products":
        chart.x_axis.axPos = "l"
        chart.y_axis.axPos = "b"
        chart.x_axis.scaling.orientation = "maxMin"
        chart.y_axis.crosses = "max"
    chart.title = plan.title
    chart.y_axis.title = plan.axis_title
    chart.x_axis.title = {"monthly": "Month", "branches": "Branch", "products": "Product"}[
        plan.kind
    ]
    chart.y_axis.numFmt = "#,##0;[Red](#,##0);0"
    chart.style = 13
    chart.height = 8
    chart.width = 18
    chart.legend = None
    values = Reference(
        sheet, min_col=plan.value_column, min_row=plan.first_row, max_row=plan.last_row
    )
    categories = Reference(
        sheet, min_col=plan.category_column, min_row=plan.first_row, max_row=plan.last_row
    )
    chart.add_data(values, titles_from_data=False)
    chart.series[0].cat = AxDataSource(strRef=StrRef(f=str(categories)))
    chart.series[0].tx = SeriesLabel(v=f"{plan.currency} net revenue (minor units)")
    if plan.kind == "monthly":
        chart.series[0].graphicalProperties.line.solidFill = "376A90"
        chart.series[0].marker.symbol = "circle"
        chart.series[0].marker.size = 5
    else:
        chart.series[0].graphicalProperties.solidFill = "376A90"
    sheet.add_chart(chart, plan.anchor)


def _render_sheet(sheet, plan: SheetPlan) -> None:
    if len(plan.rows) > EXCEL_MAX_ROWS:
        raise ReportError("Report exceeds Excel's worksheet row limit")
    for row_number, row in enumerate(plan.rows, 1):
        for column, value in enumerate(row, 1):
            _put(sheet, row_number, column, value)
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.sheet_properties.outlinePr.summaryRight = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.merge_cells("A1:G1")
    sheet.merge_cells("A2:G2")
    sheet.merge_cells("A3:G3")
    sheet["A1"].font = Font(name="Calibri", size=20, bold=True, color=_NAVY)
    sheet["A2"].font = Font(name="Calibri", size=10, color=_INK)
    sheet["A3"].font = Font(name="Calibri", size=10, italic=True, color="596D7F")
    sheet.row_dimensions[1].height = 30
    sheet.row_dimensions[2].height = 30
    sheet.row_dimensions[3].height = 26
    for row_number in plan.section_rows:
        sheet.cell(row_number, 1).font = Font(name="Calibri", size=12, bold=True, color=_NAVY)
        sheet.row_dimensions[row_number].height = 25
    for table_plan in plan.tables:
        end_column = get_column_letter(len(table_plan.headers))
        reference = f"A{table_plan.header_row}:{end_column}{table_plan.end_row}"
        for cell in sheet[table_plan.header_row][: len(table_plan.headers)]:
            cell.fill = PatternFill("solid", fgColor=_NAVY)
            cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        sheet.row_dimensions[table_plan.header_row].height = 32
        if table_plan.end_row > table_plan.header_row:
            table = Table(displayName=table_plan.name, ref=reference)
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False
            )
            sheet.add_table(table)
    primary = plan.tables[-1] if plan.title == "Data Quality" else plan.tables[0]
    sheet.auto_filter.ref = (
        f"A{primary.header_row}:{get_column_letter(len(primary.headers))}{primary.end_row}"
    )
    sheet.freeze_panes = f"A{plan.tables[0].header_row + 1}"
    sheet.print_title_rows = f"1:{primary.header_row}" if primary.header_row < 10 else "1:3"
    max_columns = max(len(row) for row in plan.rows)
    for column in range(1, max_columns + 1):
        samples = [
            text
            for row in plan.rows[:200]
            if len(row) >= column and isinstance(text := row[column - 1], str)
        ]
        width = min(36, max(16, max((min(len(s), 36) for s in samples), default=16) + 2))
        sheet.column_dimensions[get_column_letter(column)].width = width
    # Context and diagnostic JSON stay visible via wrapping rather than truncation.
    if plan.title == "Summary":
        sheet.column_dimensions["A"].width = 32
        sheet.column_dimensions["B"].width = 64
        context = next(table for table in plan.tables if table.name == "SummaryContext")
        for number in range(context.header_row + 1, context.end_row + 1):
            length = len(str(sheet.cell(number, 2).value or ""))
            sheet.row_dimensions[number].height = min(90, 16 * max(1, (length + 59) // 60))
    for table_plan in plan.tables:
        for offset, header in enumerate(table_plan.headers, 1):
            if any(term in header for term in ("amount", "revenue", "average", "price")):
                for row in sheet.iter_rows(
                    min_row=table_plan.header_row + 1,
                    max_row=table_plan.end_row,
                    min_col=offset,
                    max_col=offset,
                ):
                    row[0].alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
    for chart_plan in plan.charts:
        _render_chart(sheet, chart_plan)


def _restore_literal_text(path: Path, plan: WorkbookPlan) -> None:
    """Preserve CR/LF text on Windows XML writers without editing cell semantics."""
    # Rebuild the ZIP in memory, then rewrite only our owned staging file. XML
    # newline normalization must not change raw imported strings such as CRLF IDs.
    buffer = io.BytesIO()
    with ZipFile(path) as source, ZipFile(buffer, "w", compression=ZIP_DEFLATED) as target:
        for member in source.infolist():
            data = source.read(member.filename)
            for index, sheet in enumerate(plan.sheets, 1):
                if member.filename != f"xl/worksheets/sheet{index}.xml":
                    continue
                root = ET.fromstring(data)
                for cell in root.findall(f"{{{_MAIN}}}sheetData/{{{_MAIN}}}row/{{{_MAIN}}}c"):
                    inline = cell.find(f"{{{_MAIN}}}is/{{{_MAIN}}}t")
                    if inline is not None:
                        row, column = coordinate_to_tuple(cell.attrib["r"])
                        value = sheet.rows[row - 1][column - 1]
                        if isinstance(value, str):
                            inline.text = value
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                data = data.replace(b"\r", b"&#13;")
            target.writestr(member, data)
    path.write_bytes(buffer.getvalue())


def write_workbook(plan: WorkbookPlan, staging_path: str | Path) -> None:
    """Render a prepared plan to a caller-owned staging file; not publication."""
    if tuple(sheet.title for sheet in plan.sheets) != sheet_names(plan.mode):
        raise ReportError("Workbook must contain exactly the sheets specified for its mode")
    workbook = Workbook()
    workbook.remove(workbook.active)
    try:
        for sheet_plan in plan.sheets:
            _render_sheet(workbook.create_sheet(sheet_plan.title), sheet_plan)
        workbook.save(staging_path)
    finally:
        workbook.close()
    _restore_literal_text(Path(staging_path), plan)


def _chart_text(title) -> str:
    return (
        ""
        if title is None
        else "".join(
            node.text or "" for node in title.to_tree().iter() if node.tag.rsplit("}", 1)[-1] == "t"
        )
    )


def _verify_charts(sheet, plan: SheetPlan) -> None:
    if len(sheet._charts) != len(plan.charts):
        raise ReportError(f"{plan.title}: chart count verification failed")
    for chart, expected in zip(sheet._charts, plan.charts, strict=True):
        required_type = LineChart if expected.kind == "monthly" else BarChart
        if not isinstance(chart, required_type) or len(chart.series) != 1:
            raise ReportError(f"{plan.title}: unexpected chart type or series count")
        if expected.kind == "branches" and chart.type != "col":
            raise ReportError("Branch charts must use columns")
        if expected.kind == "products" and chart.type != "bar":
            raise ReportError("Product charts must use horizontal bars")
        values = str(
            Reference(
                sheet,
                min_col=expected.value_column,
                min_row=expected.first_row,
                max_row=expected.last_row,
            )
        )
        categories = str(
            Reference(
                sheet,
                min_col=expected.category_column,
                min_row=expected.first_row,
                max_row=expected.last_row,
            )
        )
        series = chart.series[0]
        if (
            series.val is None
            or series.val.numRef is None
            or series.val.numRef.f != values
            or series.val.numRef.numCache is not None
            or series.cat is None
            or series.cat.strRef is None
            or series.cat.strRef.f != categories
            or series.cat.strRef.strCache is not None
            or _chart_text(chart.title) != expected.title
            or _chart_text(chart.y_axis.title) != expected.axis_title
            or series.tx is None
            or series.tx.v != f"{expected.currency} net revenue (minor units)"
        ):
            raise ReportError(f"{plan.title}: chart source/currency verification failed")
        for row in range(expected.first_row, expected.last_row + 1):
            currency = sheet.cell(row, expected.currency_column).value
            value = sheet.cell(row, expected.value_column).value
            if (
                currency != expected.currency
                or type(value) is not int
                or abs(value) > CHART_INTEGER_LIMIT
            ):
                raise ReportError(f"{plan.title}: chart includes unsafe or mixed-currency data")


def verify_workbook(path: str | Path, plan: WorkbookPlan) -> None:
    """Reopen and compare every planned cell/type, table, sheet, and chart source."""
    # Retain embedded link declarations so the check below can detect them.
    # openpyxl parses their package metadata; it does not open external targets.
    workbook = load_workbook(path, data_only=False, keep_links=True)
    try:
        if tuple(workbook.sheetnames) != sheet_names(plan.mode):
            raise ReportError("Workbook sheet names/order verification failed")
        if workbook._external_links:
            raise ReportError("Report must not contain external workbook links")
        for expected in plan.sheets:
            sheet = workbook[expected.title]
            if sheet.sheet_state != "visible" or sheet.max_row != len(expected.rows):
                raise ReportError(f"{expected.title}: visibility or row count verification failed")
            width = max(7, max(len(row) for row in expected.rows))
            if sheet.max_column != width:
                raise ReportError(f"{expected.title}: column count verification failed")
            for number, cells in enumerate(sheet.iter_rows(), 1):
                record = expected.rows[number - 1]
                for column, cell in enumerate(cells):
                    value = record[column] if column < len(record) else None
                    if cell.data_type == "f":
                        raise ReportError(f"{expected.title}: unexpected formula cell")
                    normalized = "" if cell.value is None and value == "" else cell.value
                    if normalized != value or type(normalized) is not type(value):
                        raise ReportError(
                            f"{expected.title}: cell {cell.coordinate} value verification failed"
                        )
                    if isinstance(value, str) and cell.data_type not in ("s", "inlineStr"):
                        raise ReportError(f"{expected.title}: nonliteral text cell")
                    if type(value) is int and cell.data_type != "n":
                        raise ReportError(f"{expected.title}: inferred numeric type mismatch")
            expected_tables = {
                table.name: (
                    f"A{table.header_row}:{get_column_letter(len(table.headers))}{table.end_row}"
                )
                for table in expected.tables
                if table.end_row > table.header_row
            }
            if {table.name: table.ref for table in sheet.tables.values()} != expected_tables:
                raise ReportError(f"{expected.title}: table verification failed")
            primary = (
                expected.tables[-1] if expected.title == "Data Quality" else expected.tables[0]
            )
            filter_ref = (
                f"A{primary.header_row}:{get_column_letter(len(primary.headers))}{primary.end_row}"
            )
            if (
                sheet.freeze_panes != f"A{expected.tables[0].header_row + 1}"
                or sheet.auto_filter.ref != filter_ref
            ):
                raise ReportError(f"{expected.title}: filter/frozen header verification failed")
            _verify_charts(sheet, expected)
    finally:
        workbook.close()
