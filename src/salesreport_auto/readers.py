"""Strict read-only CSV and text-XLSX transport, with bounded snapshots."""

import csv
import hashlib
import io
import os
import posixpath
import stat
from collections.abc import Sequence
from pathlib import Path
from threading import RLock
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZIP_STORED, BadZipFile, ZipFile
from zlib import error as ZlibError

from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string
from openpyxl.utils.exceptions import InvalidFileException

from .config import InputLimits
from .errors import ConfigurationError, InputError
from .models import (
    ALL_COLUMNS,
    BUSINESS_COLUMNS,
    PRODUCT_ALL_COLUMNS,
    PRODUCT_COLUMNS,
    ImportRef,
    RawCell,
    RawRecord,
    ReportMode,
    TransactionInput,
)

_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_CSV_READER_LOCK = RLock()


def _snapshot(path: Path, limits: InputLimits) -> bytes:
    for component in (path, *path.parents):
        if component.is_symlink() or component.is_junction():
            raise InputError(f"{path}: symbolic links and junctions are not supported")
    with path.open("rb") as stream:
        before = os.fstat(stream.fileno())
        if not stat.S_ISREG(before.st_mode):
            raise InputError(f"{path}: input must be a regular file")
        if before.st_size > limits.max_file_bytes:
            raise InputError(f"{path}: max_file_bytes exceeded")
        data = stream.read(limits.max_file_bytes + 1)
        after = os.fstat(stream.fileno())
    current = path.stat()
    signatures = {
        (item.st_dev, item.st_ino, item.st_size, item.st_mtime_ns)
        for item in (before, after, current)
    }
    if len(data) > limits.max_file_bytes:
        raise InputError(f"{path}: max_file_bytes exceeded")
    if len(signatures) != 1 or len(data) != before.st_size:
        raise InputError(f"{path}: input changed while reading")
    return data


def _headers(values: list[object], mode: ReportMode) -> tuple[str, ...]:
    if not values or any(not isinstance(v, str) or not v.strip() for v in values):
        raise InputError("Header cells must be nonempty text")
    if len(set(values)) != len(values):
        raise InputError("Duplicate headers are not allowed")
    required = PRODUCT_COLUMNS if mode == "products" else BUSINESS_COLUMNS
    allowed = PRODUCT_ALL_COLUMNS if mode == "products" else ALL_COLUMNS
    missing = set(required) - set(values)
    unknown = set(values) - set(allowed)
    if missing or unknown:
        raise InputError(f"Header mismatch: missing={sorted(missing)}, unknown={sorted(unknown)}")
    return tuple(values)


def _read_csv(data: bytes, name: str, limits: InputLimits, mode: ReportMode) -> TransactionInput:
    # csv has a process-wide token limit smaller than the file budget. Scope and
    # restore it so valid large text identifiers do not gain an accidental limit.
    with _CSV_READER_LOCK:
        previous = csv.field_size_limit()
        try:
            csv.field_size_limit(max(previous, limits.max_file_bytes))
            return _parse_csv(data, name, limits, mode)
        finally:
            csv.field_size_limit(previous)


def _parse_csv(data: bytes, name: str, limits: InputLimits, mode: ReportMode) -> TransactionInput:
    reader = csv.reader(io.StringIO(data.decode("utf-8-sig"), newline=""), strict=True)
    headers = _headers(next(reader, []), mode)
    records = []
    while True:
        row_number = reader.line_num + 1
        values = next(reader, None)
        if values is None:
            break
        if len(records) >= limits.max_records:
            raise InputError(f"row {row_number}: max_records exceeded")
        if not values:
            values = [""] * len(headers)
        if len(values) != len(headers):
            raise InputError(f"row {row_number}: expected {len(headers)} fields, got {len(values)}")
        records.append(
            RawRecord(
                ImportRef(name, "", row_number),
                {header: RawCell(value) for header, value in zip(headers, values, strict=True)},
            )
        )
    return TransactionInput(Path(name), "", headers, "", (), tuple(records), mode)


def _check_archive(archive: ZipFile, limits: InputLimits) -> None:
    seen = set()
    total = 0
    for member in archive.infolist():
        name = member.filename
        if name in seen:
            raise InputError("Duplicate XLSX ZIP member")
        seen.add(name)
        if name.startswith("/") or "\\" in name or ".." in name.split("/"):
            raise InputError("Unsafe XLSX ZIP member path")
        if member.flag_bits & 1 or member.compress_type not in (ZIP_STORED, ZIP_DEFLATED):
            raise InputError("Encrypted or unsupported XLSX ZIP compression")
        total += member.file_size
        if total > limits.max_xlsx_uncompressed_bytes:
            raise InputError("max_xlsx_uncompressed_bytes exceeded")


def _worksheet_path(archive: ZipFile, selected: str) -> str:
    book = ET.fromstring(archive.read("xl/workbook.xml"))
    rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    sheet = next(
        item
        for item in book.findall(f"{{{_MAIN}}}sheets/{{{_MAIN}}}sheet")
        if item.attrib["name"] == selected
    )
    relationship = next(item for item in rels if item.attrib["Id"] == sheet.attrib[f"{{{_REL}}}id"])
    if relationship.attrib.get("TargetMode") == "External":
        raise InputError("Selected worksheet cannot be external")
    target = relationship.attrib["Target"]
    resolved = posixpath.normpath(
        target.lstrip("/") if target.startswith("/") else posixpath.join("xl", target)
    )
    if not resolved.startswith("xl/"):
        raise InputError("Selected worksheet escaped the workbook")
    return resolved


def _worksheet_extent(
    archive: ZipFile,
    member: str,
    limits: InputLimits,
    max_columns: int,
    shared_strings: Sequence[str],
) -> tuple[int, dict[int, int]]:
    """Inspect stored values before expanding sparse rows or stale dimensions."""
    last_value_row = 0
    nonempty_columns: dict[int, int] = {}
    previous_row = 0
    with archive.open(member) as stream:
        for _, row in ET.iterparse(stream, events=("end",)):
            if row.tag != f"{{{_MAIN}}}row":
                continue
            number = int(row.attrib.get("r", previous_row + 1))
            if not previous_row < number <= 1_048_576:
                raise InputError("Worksheet rows are out of order or out of bounds")
            previous_row = number
            previous_column = 0
            for cell in row.findall(f"{{{_MAIN}}}c"):
                if "r" in cell.attrib:
                    column_text, cell_row = coordinate_from_string(cell.attrib["r"])
                    column = column_index_from_string(column_text)
                    if cell_row != number:
                        raise InputError("Worksheet cell/row coordinate mismatch")
                else:
                    column = previous_column + 1
                if not previous_column < column <= 16_384:
                    raise InputError("Worksheet columns are out of order or out of bounds")
                previous_column = column
                stored = cell.find(f"{{{_MAIN}}}v")
                inline = cell.find(f"{{{_MAIN}}}is")
                formula = cell.find(f"{{{_MAIN}}}f")
                # XML indentation outside text nodes is not cell content.
                inline_nonempty = inline is not None and any(
                    node.text
                    for query in (f"{{{_MAIN}}}t", f"{{{_MAIN}}}r/{{{_MAIN}}}t")
                    for node in inline.findall(query)
                )
                has_value = (
                    formula is not None
                    or stored is not None
                    and stored.text is not None
                    or inline_nonempty
                    or cell.attrib.get("t") == "inlineStr"
                )
                if has_value:
                    if number - 1 > limits.max_records:
                        raise InputError(f"row {number}: max_records worksheet span exceeded")
                    last_value_row = number
                    nonempty = (
                        formula is not None
                        or stored is not None and stored.text is not None
                        or inline_nonempty
                    )
                    if cell.attrib.get("t") == "s" and stored is not None and stored.text:
                        index = int(stored.text)
                        if not 0 <= index < len(shared_strings):
                            raise InputError("Invalid shared-string reference")
                        nonempty = formula is not None or bool(shared_strings[index])
                    if nonempty or number == 1:
                        if column > max_columns:
                            raise InputError(f"row {number}: data extends beyond supported headers")
                        nonempty_columns[number] = max(nonempty_columns.get(number, 0), column)
            row.clear()
    return last_value_row, nonempty_columns


def _cell_value(cell: object) -> RawCell:
    value = cell.value
    if value is None:
        # openpyxl exposes an explicitly stored empty inline string as None.
        # Retain its distinction from an absent cell and formatting-only cells.
        if cell.data_type == "inlineStr":
            return RawCell("")
        return RawCell(None)
    kind = {
        "s": "text",
        "str": "text",
        "inlineStr": "text",
        "n": "numeric",
        "d": "date",
        "b": "boolean",
        "f": "formula",
        "e": "error",
    }.get(cell.data_type, "other")
    return RawCell(value, kind)


def _read_xlsx(data: bytes, name: str, limits: InputLimits, mode: ReportMode) -> TransactionInput:
    with ZipFile(io.BytesIO(data)) as archive:
        _check_archive(archive, limits)
        workbook = load_workbook(
            io.BytesIO(data), read_only=True, data_only=False, keep_links=False
        )
        try:
            names = tuple(sheet.title for sheet in workbook.worksheets)
            if mode == "products" and len(names) != 1:
                raise InputError("Product XLSX input requires exactly one worksheet")
            if mode == "transactions" and "Master" in names:
                selected = "Master"
            elif len(names) == 1:
                selected = names[0]
            else:
                raise InputError("XLSX requires a Master worksheet or exactly one worksheet")
            max_columns = len(PRODUCT_ALL_COLUMNS if mode == "products" else ALL_COLUMNS)
            sheet = workbook[selected]
            extent, columns = _worksheet_extent(
                archive,
                _worksheet_path(archive, selected),
                limits,
                max_columns,
                # ReadOnlyWorksheet already has the decoded shared-string table.
                # Only nonempty decoded values are forbidden beyond the headers.
                sheet._shared_strings,
            )
            sheet.reset_dimensions()
            header_cells = list(next(sheet.iter_rows(min_row=1, max_row=1, max_col=max_columns)))
            while (
                header_cells
                and header_cells[-1].value is None
                and header_cells[-1].data_type != "inlineStr"
            ):
                header_cells.pop()
            if any(_cell_value(cell).kind != "text" for cell in header_cells):
                raise InputError("Header cells must be stored as text")
            headers = _headers([cell.value for cell in header_cells], mode)
            if any(column > len(headers) for column in columns.values()):
                raise InputError("Data extends beyond source headers")
            records = []
            if extent >= 2:
                for number, cells in enumerate(
                    sheet.iter_rows(min_row=2, max_row=extent, max_col=len(headers)), 2
                ):
                    records.append(
                        RawRecord(
                            ImportRef(name, selected, number),
                            {
                                key: _cell_value(cell)
                                for key, cell in zip(headers, cells, strict=True)
                            },
                        )
                    )
            return TransactionInput(
                Path(name),
                "",
                headers,
                selected,
                tuple(sorted(set(names) - {selected})),
                tuple(records),
                mode,
            )
        finally:
            workbook.close()


def read_sales_input(
    path: str | Path, *, mode: ReportMode, limits: InputLimits | None = None
) -> TransactionInput:
    """Read one explicit-mode sales file without writing or inferring business types.

    Raises InputError on structural failures; row-level invalid values are retained
    for validation. Unsupported XLSX values retain their decoded value and kind.
    """
    if mode not in ("transactions", "products"):
        raise ConfigurationError("Mode must be transactions or products")
    target = Path(path).absolute()
    budgets = limits or InputLimits()
    if target.suffix.lower() not in (".csv", ".xlsx"):
        raise InputError(f"{target}: expected one .csv or .xlsx file")
    try:
        data = _snapshot(target, budgets)
        loader = _read_csv if target.suffix.lower() == ".csv" else _read_xlsx
        parsed = loader(data, target.name, budgets, mode)
        return TransactionInput(
            target,
            hashlib.sha256(data).hexdigest(),
            parsed.headers,
            parsed.selected_sheet,
            parsed.ignored_sheets,
            parsed.records,
            mode,
        )
    except Exception as exc:
        # Normalize parser, filesystem, ZIP/XML and openpyxl structural failures.
        # Never include offending business cell values in public diagnostics.
        if isinstance(exc, InputError):
            if str(exc).startswith(f"{target}:"):
                raise
            raise InputError(f"{target}: {exc}") from exc
        if isinstance(exc, (OSError, UnicodeError, csv.Error, ValueError, KeyError, TypeError)):
            raise InputError(
                f"{target}: cannot read {target.suffix.lower()} input ({type(exc).__name__})"
            ) from exc
        if isinstance(
            exc,
            (
                ET.ParseError,
                RuntimeError,
                StopIteration,
                IndexError,
                BadZipFile,
                EOFError,
                InvalidFileException,
                ZlibError,
                OverflowError,
                NotImplementedError,
            ),
        ):
            raise InputError(f"{target}: malformed XLSX input ({type(exc).__name__})") from exc
        raise


def read_transactions(path: str | Path, *, limits: InputLimits | None = None) -> TransactionInput:
    """Read the six-column transaction schema with its optional provenance."""
    return read_sales_input(path, mode="transactions", limits=limits)


def read_products(path: str | Path, *, limits: InputLimits | None = None) -> TransactionInput:
    """Read the ten-column product-line schema without inferring business types."""
    return read_sales_input(path, mode="products", limits=limits)
