import json
import os
import subprocess
import sys
from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from salesreport_auto import (
    PRODUCT_COLUMNS,
    InputError,
    IntegrityError,
    ReportError,
    analyze_products,
    build_report_plan,
    generate_report,
    publication,
    publish_report,
    verify_workbook,
)
from salesreport_auto.cli import main
from salesreport_auto.report_model import PRODUCT_SHEET_NAMES, SHEET_NAMES


def _table(book, sheet, name):
    area = range_boundaries(book[sheet].tables[name].ref)
    rows = list(
        book[sheet].iter_rows(
            min_col=area[0], min_row=area[1], max_col=area[2], max_row=area[3], values_only=True
        )
    )
    return [dict(zip(rows[0], row, strict=True)) for row in rows[1:]]


def _arguments(source, output):
    return ["--mode", "products", "--input", str(source), "--output", str(output)]


def test_product_workbook_has_exact_sheets_metrics_and_derived_raw_fields(
    write_products, product_line, transport, tmp_path
):
    source = write_products(
        [
            product_line(transaction_id="T1"),
            product_line(
                transaction_id="T1",
                line_id="L2",
                product_id="P2",
                quantity="1",
                unit_price="4.00",
                amount="4.00",
            ),
            product_line(transaction_id="T2", quantity="-1", amount="-3.00"),
        ],
        format=transport,
    )
    before = source.read_bytes(), source.stat().st_mtime_ns
    report = generate_report(source, tmp_path / "products.xlsx", mode="products")
    assert report.status == "completed" and report.plan.mode == "products"
    assert report.to_summary()["mode"] == "products"
    assert report.to_summary()["local_counts"] == {
        "input": 3,
        "accepted": 3,
        "rejected": 0,
        "redundant": 0,
    }
    assert report.to_summary()["currencies"] == [
        {"code": "USD", "precision": 2, "net_revenue": "7.00", "transaction_count": 2}
    ]
    book = load_workbook(report.output, data_only=False)
    try:
        assert tuple(book.sheetnames) == PRODUCT_SHEET_NAMES
        assert all(sheet.sheet_state == "visible" for sheet in book)
        summary = _table(book, "Summary", "SummaryMetrics")[0]
        assert summary["transaction_count"] == 2 and summary["accepted_record_count"] == 3
        assert summary["average_transaction_value"] == "3.50"
        products = _table(book, "Products", "ProductMetrics")
        assert products == [
            {
                "currency": "USD",
                "revenue_rank": 1,
                "product_id": "P2",
                "product_name": "P2",
                "accepted_line_count": 1,
                "transaction_count": 1,
                "unique_customer_count": 1,
                "positive_quantity": 1,
                "negative_quantity_magnitude": 0,
                "net_quantity": 1,
                "net_revenue": "4.00",
                "minimum_unit_price": "4.00",
                "maximum_unit_price": "4.00",
                "positive_quantity_weighted_unit_price": "4.00",
            },
            {
                "currency": "USD",
                "revenue_rank": 2,
                "product_id": "P1",
                "product_name": "P1",
                "accepted_line_count": 2,
                "transaction_count": 2,
                "unique_customer_count": 1,
                "positive_quantity": 2,
                "negative_quantity_magnitude": 1,
                "net_quantity": 1,
                "net_revenue": "3.00",
                "minimum_unit_price": "3.00",
                "maximum_unit_price": "3.00",
                "positive_quantity_weighted_unit_price": "3.00",
            },
        ]
        assert book["Products"]["K7"].data_type == "s"
        assert book["Products"]["L7"].data_type == "s"
        assert book["Products"]["N7"].data_type == "s"
        raw = _table(book, "Raw Data", "ImportedRecords")
        assert raw[2]["quantity"] == "-1" and raw[2]["report_quantity"] == -1
        assert raw[2]["unit_price"] == "3.00" and raw[2]["report_unit_price_minor"] == "300"
        assert raw[2]["report_amount_minor"] == "-300"
        assert list(raw[0]) == [
            *PRODUCT_COLUMNS,
            "report_import_file",
            "report_import_sheet",
            "report_import_row",
            "report_status",
            "report_reason_codes",
            "report_issues",
            "report_warnings",
            "report_duplicate_of",
            "report_amount_minor",
            "report_currency_precision",
            "report_normalized_date",
            "report_raw_cell_kinds",
            "report_quantity",
            "report_unit_price_minor",
        ]
        context = {row["field"]: row["value"] for row in _table(book, "Summary", "SummaryContext")}
        assert context["mode"] == "products"
        assert context["record_grain"] == "Product line occurrences"
        assert context["upstream_status"] == "Unknown / not supplied"
        assert len(book["Products"]._charts) == 1
        assert book["Products"]._charts[0].type == "bar"
    finally:
        book.close()
    assert before == (source.read_bytes(), source.stat().st_mtime_ns)
    verify_workbook(report.output, report.plan)


def test_product_chart_is_top_ten_per_currency_with_isolated_integer_helpers(
    write_products, product_line, tmp_path
):
    rows = [
        product_line(
            transaction_id=f"{currency}-{number}",
            product_id=f"P{number:02}",
            quantity="1",
            unit_price=str(number),
            amount=str(number),
            currency=currency,
        )
        for currency in ("USD", "PLN")
        for number in range(12)
    ]
    report = generate_report(write_products(rows), tmp_path / "ranked.xlsx", mode="products")
    book = load_workbook(report.output)
    try:
        products = book["Products"]
        assert len(_table(book, "Products", "ProductMetrics")) == 24
        assert len(products._charts) == 2
        for chart, currency in zip(products._charts, ("PLN", "USD"), strict=True):
            assert chart.type == "bar" and len(chart.series) == 1
            series = chart.series[0]
            assert series.tx.v.startswith(currency)
            assert series.val.numRef.f.startswith("'Products'!")
            assert series.cat.strRef.f.startswith("'Products'!")
            value_area = range_boundaries(series.val.numRef.f.split("!")[1])
            category_area = range_boundaries(series.cat.strRef.f.split("!")[1])
            assert value_area[0] == value_area[2] == 19
            assert category_area[0] == category_area[2] == 18
            assert value_area[3] - value_area[1] + 1 == 10
            assert [
                products.cell(row, 18).value for row in range(value_area[1], value_area[3] + 1)
            ] == [f"P{n:02}" for n in range(11, 1, -1)]
            assert all(
                products.cell(row, 17).value == currency
                for row in range(value_area[1], value_area[3] + 1)
            )
            assert all(
                type(products.cell(row, 19).value) is int
                for row in range(value_area[1], value_area[3] + 1)
            )
    finally:
        book.close()


@pytest.mark.parametrize("plotted", [True, False])
def test_product_chart_limit_applies_only_to_plotted_ranks(
    write_products, product_line, tmp_path, plotted
):
    rows = [
        product_line(
            transaction_id=str(i),
            product_id=f"P{i}",
            quantity="1",
            unit_price="1",
            amount="1",
            currency="JPY",
        )
        for i in range(10)
    ]
    rows.append(
        product_line(
            transaction_id="big",
            product_id="big",
            quantity="1" if plotted else "-1",
            unit_price="1000000000000000",
            amount="1000000000000000" if plotted else "-1000000000000000",
            currency="JPY",
        )
    )
    report = generate_report(write_products(rows), tmp_path / "precision.xlsx", mode="products")
    assert ("Products: JPY" in report.plan.omitted_charts) == plotted
    book = load_workbook(report.output)
    try:
        assert len(book["Products"]._charts) == int(not plotted)
        product = next(
            row for row in _table(book, "Products", "ProductMetrics") if row["product_id"] == "big"
        )
        assert product["net_revenue"] == ("1000000000000000" if plotted else "-1000000000000000")
    finally:
        book.close()


def test_formula_like_product_ids_names_and_original_quantity_are_literal(
    write_products, product_line, tmp_path
):
    source = write_products(
        [
            product_line(
                line_id="+001",
                product_id="=商品\r\n+line",
                product_name="@Label",
                quantity="+0002",
                unit_price="+03.000",
            )
        ],
        headers=[*PRODUCT_COLUMNS, "product_name"],
    )
    report = generate_report(source, tmp_path / "literal-products.xlsx", mode="products")
    book = load_workbook(report.output, data_only=False)
    try:
        assert book["Products"]["C7"].value == "=商品\r\n+line"
        assert book["Products"]["C7"].data_type == "s"
        assert book["Products"]["D7"].value == "@Label"
        assert book["Products"]["D7"].data_type == "s"
        raw = _table(book, "Raw Data", "ImportedRecords")[0]
        assert (raw["line_id"], raw["quantity"], raw["unit_price"]) == ("+001", "+0002", "+03.000")
        assert (raw["report_quantity"], raw["report_unit_price_minor"]) == (2, "300")
        for sheet in book:
            assert all(cell.data_type != "f" for row in sheet for cell in row)
    finally:
        book.close()


def test_product_label_warnings_redundancy_and_quarantine_are_fully_disclosed(
    write_products, product_line, tmp_path
):
    rows = [
        product_line(transaction_id="A", product_name="Zulu"),
        product_line(transaction_id="A", product_name="Zulu"),
        product_line(transaction_id="B", product_name="Alpha"),
        product_line(transaction_id="bad", product_name="Other"),
        product_line(
            transaction_id="bad", line_id="L2", quantity="not-a-quantity", product_name="Other"
        ),
    ]
    report = generate_report(
        write_products(rows, headers=[*PRODUCT_COLUMNS, "product_name"]),
        tmp_path / "quality.xlsx",
        mode="products",
    )
    assert report.review_reasons == ("LOCAL_REJECTIONS", "PRODUCT_LABEL_CONFLICT")
    assert report.analysis.quality.local_monetary_complete
    book = load_workbook(report.output)
    try:
        assert _table(book, "Products", "ProductMetrics")[0]["product_name"] == "Alpha"
        details = _table(book, "Data Quality", "QualityDetails")
        assert len(details) == 5  # Two accepted warnings, one duplicate, two rejections.
        assert sum(row["status"] == "rejected" for row in details) == 2
        assert sum("PRODUCT_LABEL_CONFLICT" in row["warnings"] for row in details) == 2
        assert len(_table(book, "Raw Data", "ImportedRecords")) == 5
        assert _table(book, "Data Quality", "LocalQuality")[0]["accepted"] == 2
    finally:
        book.close()


@pytest.mark.parametrize("case", ["header_only", "rejected", "negative_only", "zero_quantity"])
def test_product_empty_and_no_positive_quantity_reports(
    write_products, product_line, tmp_path, case
):
    if case == "header_only":
        rows = []
    elif case == "rejected":
        rows = [product_line(quantity="bad")]
    elif case == "negative_only":
        rows = [product_line(quantity="-1", amount="-3.00")]
    else:
        rows = [product_line(quantity="0", amount="0.00")]
    report = generate_report(write_products(rows), tmp_path / "empty.xlsx", mode="products")
    book = load_workbook(report.output)
    try:
        assert tuple(book.sheetnames) == PRODUCT_SHEET_NAMES
        assert book["Products"]["A6"].value == "currency"
        if case in ("header_only", "rejected"):
            assert "ProductMetrics" not in book["Products"].tables
            assert not book["Products"]._charts
            assert report.review_reasons[0] == "NO_ACCEPTED_RECORDS"
        else:
            assert (
                _table(book, "Products", "ProductMetrics")[0][
                    "positive_quantity_weighted_unit_price"
                ]
                == "N/A"
            )
    finally:
        book.close()


def test_cli_products_status_summary_and_module_entry_point(
    write_products, product_line, tmp_path, capsys
):
    source = write_products([product_line(), product_line(line_id="L2")])
    output = tmp_path / "cli-products.xlsx"
    assert main(_arguments(source, output)) == 0
    captured = capsys.readouterr()
    assert captured.err == ""
    summary = json.loads(captured.out)
    assert summary["mode"] == "products" and summary["upstream"] is None
    assert summary["local_counts"]["accepted"] == 2
    assert summary["currencies"][0]["transaction_count"] == 1
    assert summary["currencies"][0]["net_revenue"] == "12.00"
    environment = os.environ.copy()
    environment["PYTHONPATH"] = str(Path(__file__).parents[1] / "src")
    process = subprocess.run(
        [
            sys.executable,
            "-m",
            "salesreport_auto",
            *_arguments(source, tmp_path / "module-products.xlsx"),
        ],
        env=environment,
        capture_output=True,
        text=True,
        check=False,
    )
    assert process.returncode == 0 and process.stderr == ""
    assert json.loads(process.stdout)["mode"] == "products"


@pytest.mark.parametrize(
    "case,exit_code", [("invalid", 1), ("label_warning", 1), ("wrong_schema", 2)]
)
def test_cli_product_review_and_schema_errors(
    write_products, write_input, product_line, record, tmp_path, capsys, case, exit_code
):
    if case == "invalid":
        source = write_products([product_line(quantity="bad")])
    elif case == "label_warning":
        source = write_products(
            [product_line(product_name="A"), product_line(line_id="L2", product_name="B")],
            headers=[*PRODUCT_COLUMNS, "product_name"],
        )
    else:
        source = write_input([record()])
    output = tmp_path / "review.xlsx"
    assert main(_arguments(source, output)) == exit_code
    captured = capsys.readouterr()
    if exit_code == 1:
        assert json.loads(captured.out)["status"] == "review_required"
        assert captured.err == "" and output.exists()
    else:
        assert captured.out == "" and "Header mismatch" in captured.err
        assert not output.exists()


def test_transaction_mode_still_has_five_sheets_and_rejects_product_schema(
    write_input, write_products, record, product_line, tmp_path
):
    report = generate_report(write_input([record()]), tmp_path / "transactions.xlsx")
    assert tuple(sheet.title for sheet in report.plan.sheets) == SHEET_NAMES
    assert report.to_summary()["mode"] == "transactions"
    with pytest.raises(InputError):
        generate_report(write_products([product_line()]), tmp_path / "wrong.xlsx")


def test_product_chart_corruption_and_stale_aggregates_are_rejected(
    write_products, product_line, tmp_path, monkeypatch
):
    source = write_products([product_line()])
    analysis = analyze_products(source)
    stale = replace(analysis, aggregates=replace(analysis.aggregates, products=()))
    with pytest.raises(IntegrityError, match="aggregates"):
        build_report_plan(stale, application_version="test")
    real_write = publication.write_workbook

    def corrupt(plan, staging):
        real_write(plan, staging)
        book = load_workbook(staging)
        book["Products"]._charts[0].series[0].val.numRef.f = "'Branches'!$L$7"
        book.save(staging)
        book.close()

    monkeypatch.setattr(publication, "write_workbook", corrupt)
    with pytest.raises(ReportError, match="chart source"):
        publish_report(analysis, tmp_path / "corrupt.xlsx")
    assert not (tmp_path / "corrupt.xlsx").exists()
    assert not list(tmp_path.glob(".corrupt-*.xlsx"))
