"""Deterministic synthetic transport fixtures; no sibling repository is needed."""

import csv

import pytest
from openpyxl import Workbook

from salesreport_auto import BUSINESS_COLUMNS, PRODUCT_COLUMNS


def pytest_addoption(parser):
    parser.addoption(
        "--sheetmerge-integration",
        action="store_true",
        default=False,
        help="Run optional sibling SheetMerge fixture integration tests",
    )


def pytest_collection_modifyitems(config, items):
    if not config.getoption("--sheetmerge-integration"):
        reason = pytest.mark.skip(
            reason="Opt-in SheetMerge integration; use --sheetmerge-integration"
        )
        for item in items:
            if "sheetmerge_integration" in item.keywords:
                item.add_marker(reason)


@pytest.fixture
def bundle_factory(tmp_path):
    from bundle_fixtures import make_bundle

    counter = 0

    def make(**options):
        nonlocal counter
        counter += 1
        return make_bundle(tmp_path / f"bundle-{counter}", **options)

    return make


@pytest.fixture
def record():
    def make(**changes):
        row = {
            "branch_id": "001",
            "transaction_id": "000001",
            "transaction_date": "2024-01-02",
            "customer_id": "000101",
            "amount": "12.34",
            "currency": "USD",
        }
        row.update(changes)
        return row

    return make


@pytest.fixture
def write_input(tmp_path):
    def write(rows, *, format="csv", headers=None, name=None, sheet="Master", extra_sheets=()):
        headers = list(BUSINESS_COLUMNS if headers is None else headers)
        path = tmp_path / (name or f"sales.{format}")
        values = [
            [row.get(header, "") for header in headers] if isinstance(row, dict) else row
            for row in rows
        ]
        if format == "csv":
            with path.open("w", encoding="utf-8", newline="") as stream:
                writer = csv.writer(stream)
                writer.writerow(headers)
                writer.writerows(values)
        else:
            book = Workbook()
            ws = book.active
            ws.title = sheet
            for number, row in enumerate([headers, *values], 1):
                for column, value in enumerate(row, 1):
                    cell = ws.cell(number, column, value)
                    if isinstance(value, str):
                        cell.data_type = "s"
            for title in extra_sheets:
                book.create_sheet(title)
            book.save(path)
            book.close()
        return path

    return write


@pytest.fixture(params=("csv", "xlsx"))
def transport(request):
    return request.param


@pytest.fixture
def product_line(record):
    def make(**changes):
        row = record(amount="6.00")
        row.update(line_id="L1", product_id="P1", quantity="2", unit_price="3.00")
        row.update(changes)
        return row

    return make


@pytest.fixture
def write_products(write_input):
    def write(rows, *, format="csv", headers=None, **options):
        columns = list(PRODUCT_COLUMNS if headers is None else headers)
        return write_input(rows, format=format, headers=columns, **options)

    return write
