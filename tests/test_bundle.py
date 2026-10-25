import csv
import hashlib
import json
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from salesreport_auto import (
    BundleError,
    CurrencyConfig,
    ReportError,
    analyze_transactions,
    engine,
    generate_report,
    load_bundle,
    publication,
    publish_report,
    verify_workbook,
)
from salesreport_auto.cli import main


def _table(book, name, table):
    sheet = book[name]
    bounds = range_boundaries(sheet.tables[table].ref)
    rows = list(
        sheet.iter_rows(
            min_col=bounds[0],
            min_row=bounds[1],
            max_col=bounds[2],
            max_row=bounds[3],
            values_only=True,
        )
    )
    return [dict(zip(rows[0], row, strict=True)) for row in rows[1:]]


@pytest.mark.parametrize("master", ["master.csv", "master.xlsx"])
def test_bundle_report_keeps_local_and_upstream_separate(bundle_factory, tmp_path, master):
    fixture = bundle_factory()
    before = {
        p.name: (hashlib.sha256(p.read_bytes()).hexdigest(), p.stat().st_mtime_ns)
        for p in fixture.root.iterdir()
    }
    report = generate_report(fixture.root / master, tmp_path / "report.xlsx", bundle=fixture.root)
    assert report.status == "review_required"
    assert report.review_reasons == ("UPSTREAM_REJECTIONS", "UPSTREAM_UNQUANTIFIED_AMOUNTS")
    assert report.analysis.quality.counts.input == report.analysis.quality.counts.accepted == 1
    assert report.analysis.quality.counts.rejected == report.analysis.quality.counts.redundant == 0
    assert report.analysis.quality.status == "completed"
    summary = report.to_summary()
    assert summary["upstream"] == {
        "counts": {"input": 5, "accepted": 1, "rejected": 3, "redundant": 1},
        "unquantified_count": 2,
        "monetary_complete": False,
    }
    assert summary["currencies"][0]["net_revenue"] == "10.00"
    assert report.analysis.upstream.reconciliation.unknown_currency_count == 1
    assert len(report.analysis.upstream.exclusions) == 4
    assert dict(report.analysis.classified.validated.config.precision) == {"USD": 2}
    book = load_workbook(report.output)
    try:
        assert book.sheetnames == [
            "Summary",
            "Monthly Sales",
            "Branches",
            "Data Quality",
            "Raw Data",
        ]
        context = {row["field"]: row["value"] for row in _table(book, "Summary", "SummaryContext")}
        assert context["local_input_count"] == 1 and context["upstream_input_count"] == 5
        assert context["upstream_rejected_count"] == 3
        assert context["upstream_acceptance_complete"] is False
        assert context["upstream_monetary_complete"] is False
        assert "not authenticity" in context["bundle_verification_scope"]
        assert "not established" in book["Summary"]["A3"].value
        assert len(_table(book, "Raw Data", "ImportedRecords")) == 1
        assert len(_table(book, "Data Quality", "QualityDetails")) == 4
        assert all(
            row["scope"] == "upstream" for row in _table(book, "Data Quality", "QualityDetails")
        )
        assert _table(book, "Data Quality", "UpstreamLedgers")[0]["input_known_amount"] == "23.00"
        assert _table(book, "Data Quality", "LocalLedgers")[0]["input_known_amount"] == "10.00"
        assert len(_table(book, "Data Quality", "UpstreamExclusions")) == 4
        assert _table(book, "Data Quality", "UpstreamQuality")[0]["unquantified_count"] == 2
    finally:
        book.close()
    verify_workbook(report.output, report.plan)
    assert before == {
        p.name: (hashlib.sha256(p.read_bytes()).hexdigest(), p.stat().st_mtime_ns)
        for p in fixture.root.iterdir()
    }


def test_same_master_standalone_has_identical_revenue_but_unknown_upstream(
    bundle_factory, tmp_path
):
    fixture = bundle_factory()
    selected = fixture.root / "master.csv"
    bundled = analyze_transactions(selected, bundle=fixture.root)
    standalone = generate_report(selected, tmp_path / "standalone.xlsx")
    assert bundled.aggregates == standalone.analysis.aggregates
    assert standalone.analysis.upstream is None
    assert standalone.to_summary()["upstream"] is None and standalone.status == "completed"
    book = load_workbook(standalone.output)
    try:
        assert (
            _table(book, "Data Quality", "UpstreamQuality")[0]["accepted"]
            == "Unknown / not supplied"
        )
    finally:
        book.close()


def test_bundle_precision_is_complete_not_merged_and_cli_checks_only_explicit_overrides(
    bundle_factory, tmp_path, capsys
):
    fixture = bundle_factory(precision={"KWD": 3})
    source = fixture.root / "master.csv"
    result = analyze_transactions(source, bundle=fixture.root, config=CurrencyConfig({"KWD": 3}))
    assert dict(result.classified.validated.config.precision) == {"KWD": 3}
    assert dict(result.classified.validated.config.origins) == {"KWD": "bundle"}
    assert result.aggregates.summary[0].metrics.net_revenue == "10.000"
    for definition in ("KWD=2", "PLN=2"):
        output = tmp_path / f"{definition[:3]}.xlsx"
        assert (
            main(
                [
                    "--mode",
                    "transactions",
                    "--input",
                    str(source),
                    "--output",
                    str(output),
                    "--bundle",
                    str(fixture.root),
                    "--currency-precision",
                    definition,
                ]
            )
            == 2
        )
        assert not output.exists()
        captured = capsys.readouterr()
        assert captured.out == "" and "INVALID_BUNDLE" in captured.err


@pytest.mark.parametrize(
    "variant,status",
    [
        ("clean", "completed"),
        ("duplicates", "completed"),
        ("conflict", "review_required"),
        ("header", "review_required"),
    ],
)
def test_clean_duplicate_and_empty_master_bundle_statuses(
    bundle_factory, tmp_path, variant, status
):
    fixture = bundle_factory(variant=variant)
    result = generate_report(
        fixture.root / "master.csv", tmp_path / "report.xlsx", bundle=fixture.root
    )
    assert result.status == status
    if variant in ("conflict", "header"):
        assert result.review_reasons[0] == "NO_ACCEPTED_RECORDS"
        assert not result.analysis.aggregates.summary
    if variant == "conflict":
        assert result.review_reasons == ("NO_ACCEPTED_RECORDS", "UPSTREAM_REJECTIONS")
        assert result.analysis.upstream.reconciliation.monetary_complete


def test_selected_master_can_be_identical_copy_outside_bundle(bundle_factory, tmp_path):
    fixture = bundle_factory()
    copied = tmp_path / "master.csv"
    shutil.copyfile(fixture.root / "master.csv", copied)
    assert analyze_transactions(copied, bundle=fixture.root).quality.counts.accepted == 1
    wrong_name = tmp_path / "renamed.csv"
    shutil.copyfile(copied, wrong_name)
    with pytest.raises(BundleError, match="BUNDLE_MISMATCH"):
        analyze_transactions(wrong_name, bundle=fixture.root)
    copied.write_bytes(copied.read_bytes() + b"\n")
    with pytest.raises(BundleError, match="BUNDLE_MISMATCH"):
        analyze_transactions(copied, bundle=fixture.root)


@pytest.mark.parametrize(
    "member",
    [
        "master.csv",
        "master.xlsx",
        "exclusions.csv",
        "audit.jsonl",
        "reconciliation.csv",
        "reconciliation.json",
        "manifest.json",
    ],
)
def test_missing_required_member_is_fatal(bundle_factory, member):
    fixture = bundle_factory()
    (fixture.root / member).unlink()
    with pytest.raises(BundleError, match="INVALID_BUNDLE"):
        analyze_transactions(fixture.root / "master.csv", bundle=fixture.root)


@pytest.mark.parametrize(
    "member",
    [
        "master.csv",
        "master.xlsx",
        "exclusions.csv",
        "audit.jsonl",
        "reconciliation.csv",
        "reconciliation.json",
    ],
)
def test_altered_member_digest_is_fatal(bundle_factory, member):
    fixture = bundle_factory()
    path = fixture.root / member
    path.write_bytes(path.read_bytes() + b" ")
    with pytest.raises(BundleError, match="SHA-256 mismatch"):
        load_bundle(fixture.root)


@pytest.mark.parametrize(
    "field,value",
    [
        ("tool", "other"),
        ("tool_version", 1),
        ("manifest_version", 2),
        ("manifest_version", True),
        ("business_key", ["transaction_id"]),
        ("configuration_sha256", "BAD"),
        ("inputs", "not-array"),
        ("policies", {}),
        ("configuration", {}),
    ],
)
def test_manifest_requirements_fail_closed(bundle_factory, field, value):
    fixture = bundle_factory()
    manifest = fixture.manifest()
    manifest[field] = value
    fixture.write_manifest(manifest)
    with pytest.raises(BundleError, match="INVALID_BUNDLE"):
        load_bundle(fixture.root)


@pytest.mark.parametrize(
    "case",
    [
        "extra_path",
        "extra_name",
        "missing_hash",
        "bad_digest",
        "bad_precision",
        "empty_precision",
        "bad_profiles",
        "bad_limits",
        "duplicate_source",
        "source_hash",
    ],
)
def test_nested_manifest_requirements(bundle_factory, case):
    fixture = bundle_factory()
    data = fixture.manifest()
    if case == "extra_path":
        data["output_sha256"]["../outside.csv"] = "0" * 64
    elif case == "extra_name":
        data["output_sha256"]["other.csv"] = "0" * 64
    elif case == "missing_hash":
        del data["output_sha256"]["audit.jsonl"]
    elif case == "bad_digest":
        data["output_sha256"]["audit.jsonl"] = "F" * 64
    elif case == "bad_precision":
        data["configuration"]["currency_precision"]["USD"] = True
    elif case == "empty_precision":
        data["configuration"]["currency_precision"] = {}
    elif case == "bad_profiles":
        data["configuration"]["profiles"] = {}
    elif case == "bad_limits":
        data["configuration"]["limits"] = []
    elif case == "duplicate_source":
        data["inputs"].append(data["inputs"][0])
    else:
        data["inputs"][0]["sha256"] = "bad"
    fixture.write_manifest(data)
    with pytest.raises(BundleError):
        load_bundle(fixture.root)


@pytest.mark.parametrize(
    "case",
    [
        "bool_count",
        "negative_count",
        "count_equation",
        "count_balance",
        "known_flag",
        "complete_flag",
        "unknown_count",
        "unquantified_count",
        "missing_currency",
        "duplicate_currency",
        "precision",
        "amounts",
        "money_partition",
        "unquantified_partition",
        "accepted_unquantified",
        "currency_coverage",
    ],
)
def test_reconciliation_invariants_not_just_file_hashes(bundle_factory, case):
    fixture = bundle_factory()

    def mutate(data):
        row = data["currencies"][0]
        if case == "bool_count":
            data["counts"]["accepted"] = True
        elif case == "negative_count":
            data["counts"]["rejected"] = -1
        elif case == "count_equation":
            data["count_equation"] = "different"
        elif case == "count_balance":
            data["counts"]["input"] = 20
        elif case == "known_flag":
            data["known_amounts_balanced"] = False
        elif case == "complete_flag":
            data["monetary_complete"] = True
        elif case == "unknown_count":
            data["unknown_currency_count"] = 3
        elif case == "unquantified_count":
            data["unquantified_count"] = 1
        elif case == "missing_currency":
            data["currencies"] = []
        elif case == "duplicate_currency":
            data["currencies"].append(row.copy())
        elif case == "precision":
            row["precision"] = 3
        elif case == "amounts":
            row["amounts"]["input"] = "23.01"
        elif case == "money_partition":
            row["minor_units"]["input"] = "2301"
            row["amounts"]["input"] = "23.01"
        elif case == "unquantified_partition":
            row["unquantified_counts"]["input"] = 2
        elif case == "accepted_unquantified":
            row["unquantified_counts"] = {"input": 1, "accepted": 1, "rejected": 0, "redundant": 0}
        else:
            row["monetary_complete"] = True

    fixture.change_reconciliation(mutate)
    with pytest.raises(BundleError, match="INVALID_BUNDLE"):
        load_bundle(fixture.root)


def test_manifest_embedded_report_and_strict_json_are_checked(bundle_factory):
    for token in (
        '{"tool":"sheetmerge","tool":"sheetmerge"}',
        '{"x":NaN}',
        '{"x":Infinity}',
        "not-json",
    ):
        fixture = bundle_factory()
        (fixture.root / "manifest.json").write_text(token, encoding="utf-8")
        with pytest.raises(BundleError):
            load_bundle(fixture.root)
    fixture = bundle_factory()
    data = fixture.manifest()
    data["reconciliation"]["counts"]["accepted"] = True
    fixture.write_manifest(data)
    with pytest.raises(BundleError):
        load_bundle(fixture.root)
    fixture = bundle_factory()
    data = fixture.manifest()
    data["reconciliation"]["extra"] = "different"
    fixture.write_manifest(data)
    with pytest.raises(BundleError, match="disagrees"):
        load_bundle(fixture.root)


@pytest.mark.parametrize(
    "case",
    [
        "reference",
        "source",
        "duplicate_provenance",
        "status",
        "reason_codes",
        "reasons",
        "raw_values",
        "rejected_reference",
        "row_number",
        "currency",
        "count",
    ],
)
def test_exclusions_and_representative_references_are_required(bundle_factory, case):
    fixture = bundle_factory()
    rows = fixture.exclusions()
    if case == "reference":
        rows[0]["duplicate_of_row"] = "999"
    elif case == "source":
        rows[0]["source_file"] = "not-in-inventory.csv"
    elif case == "duplicate_provenance":
        rows[0]["source_row"] = "2"
    elif case == "status":
        rows[0]["status"] = "accepted"
    elif case == "reason_codes":
        rows[0]["reason_codes"] = '["OTHER"]'
    elif case == "reasons":
        rows[0]["reasons"] = '[{"code": "EXACT_DUPLICATE", "field": 1, "message":"x"}]'
    elif case == "raw_values":
        rows[0]["raw_values"] = '{"x":1,"x":2}'
    elif case == "rejected_reference":
        rows[1]["duplicate_of_row"] = "2"
    elif case == "row_number":
        rows[0]["source_row"] = "0"
    elif case == "currency":
        rows[0]["currency"] = "PLN"
    else:
        rows.pop()
    fixture.write_exclusions(rows)
    with pytest.raises(BundleError):
        analyze_transactions(fixture.root / "master.csv", bundle=fixture.root)


@pytest.mark.parametrize(
    "change", ["amount", "duplicate_key", "missing_row", "missing_provenance", "invalid_value"]
)
def test_resigned_master_still_must_agree_with_reconciliation(bundle_factory, change):
    fixture = bundle_factory(variant="clean")
    path = fixture.root / "master.csv"
    with path.open(encoding="utf-8", newline="") as stream:
        reader = csv.DictReader(stream)
        headers, rows = reader.fieldnames, list(reader)
    if change == "amount":
        rows[0]["amount"] = "11.00"
    elif change == "duplicate_key":
        rows.append({**rows[0], "source_row": "3"})
    elif change == "missing_row":
        rows = []
    elif change == "missing_provenance":
        headers = headers[:6]
        rows = [{key: row[key] for key in headers} for row in rows]
    else:
        rows[0]["transaction_date"] = "bad"
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)
    fixture.resign()
    with pytest.raises(BundleError, match="BUNDLE_MISMATCH"):
        analyze_transactions(path, bundle=fixture.root)


def test_additive_metadata_and_unknown_reason_codes_are_retained(bundle_factory):
    fixture = bundle_factory()
    data = fixture.manifest()
    data["status"] = "completed"  # Not an actual coverage field; must not suppress review.
    data["tool_version"] = "future-with-manifest-v1"
    data["configuration"]["future"] = {"notes": ["retained"]}
    fixture.write_manifest(data)
    rows = fixture.exclusions()
    reasons = json.loads(rows[1]["reasons"])
    reasons[0]["code"] = "UNKNOWN_NONEMPTY_REASON"
    rows[1]["reasons"] = json.dumps(reasons)
    rows[1]["reason_codes"] = '["UNKNOWN_NONEMPTY_REASON"]'
    fixture.write_exclusions(rows)
    analysis = analyze_transactions(fixture.root / "master.csv", bundle=fixture.root)
    assert analysis.upstream.review_reasons
    assert "UNKNOWN_NONEMPTY_REASON" in analysis.upstream.reason_counts
    assert analysis.upstream.configuration["future"]["notes"] == ("retained",)
    with pytest.raises(TypeError):
        analysis.upstream.configuration["future"] = None


def test_cli_bundle_and_products_restriction(bundle_factory, tmp_path, capsys):
    fixture = bundle_factory()
    args = [
        "--mode",
        "transactions",
        "--input",
        str(fixture.root / "master.csv"),
        "--output",
        str(tmp_path / "cli.xlsx"),
        "--bundle",
        str(fixture.root),
    ]
    assert main(args) == 1
    result = json.loads(capsys.readouterr().out)
    assert result["upstream"]["counts"]["input"] == 5
    args[1] = "products"
    assert main(args) == 2
    captured = capsys.readouterr()
    assert captured.out == "" and "only in transaction mode" in captured.err


def test_bundle_must_remain_unchanged_and_output_must_stay_outside(
    bundle_factory, tmp_path, monkeypatch
):
    fixture = bundle_factory()
    analysis = analyze_transactions(fixture.root / "master.csv", bundle=fixture.root)
    with pytest.raises(ReportError, match="outside"):
        publish_report(analysis, fixture.root / "report.xlsx")
    real_write = publication.write_workbook

    def change_bundle(plan, staging):
        real_write(plan, staging)
        with (fixture.root / "audit.jsonl").open("a", encoding="utf-8") as stream:
            stream.write("changed\n")

    monkeypatch.setattr(publication, "write_workbook", change_bundle)
    with pytest.raises(ReportError):
        publish_report(analysis, tmp_path / "changed.xlsx")
    assert not (tmp_path / "changed.xlsx").exists()
    assert not list(tmp_path.glob(".changed-*.xlsx"))


def test_cli_invalid_bundle_never_falls_back_to_standalone(bundle_factory, tmp_path, capsys):
    fixture = bundle_factory()
    (fixture.root / "audit.jsonl").write_text("altered", encoding="utf-8")
    output = tmp_path / "no-fallback.xlsx"
    assert (
        main(
            [
                "--mode",
                "transactions",
                "--input",
                str(fixture.root / "master.csv"),
                "--output",
                str(output),
                "--bundle",
                str(fixture.root),
            ]
        )
        == 2
    )
    captured = capsys.readouterr()
    assert captured.out == "" and "INVALID_BUNDLE" in captured.err
    assert not output.exists()


def test_implicit_defaults_do_not_override_custom_bundle_usd_precision(bundle_factory):
    fixture = bundle_factory(variant="clean", precision={"USD": 3, "JPY": 0})
    result = analyze_transactions(
        fixture.root / "master.csv", bundle=fixture.root, config=CurrencyConfig()
    )
    assert result.classified.validated.config.precision["USD"] == 3
    assert result.aggregates.summary[0].metrics.net_revenue == "10.000"
    assert {ledger.currency for ledger in result.upstream.reconciliation.currencies} == {
        "USD",
        "JPY",
    }
    with pytest.raises(BundleError, match="Explicit currency precision"):
        analyze_transactions(
            fixture.root / "master.csv", bundle=fixture.root, config=CurrencyConfig({"USD": 2})
        )


def test_bundle_links_are_refused_without_os_symlink_privileges(bundle_factory, monkeypatch):
    fixture = bundle_factory()
    selected = fixture.root / "audit.jsonl"
    real = Path.is_symlink
    monkeypatch.setattr(Path, "is_symlink", lambda path: path == selected or real(path))
    with pytest.raises(BundleError, match="links"):
        load_bundle(fixture.root)


def test_bundle_change_during_master_analysis_is_detected(bundle_factory, monkeypatch):
    fixture = bundle_factory()
    real_read = engine.read_transactions

    def change_after_read(*args, **kwargs):
        source = real_read(*args, **kwargs)
        (fixture.root / "audit.jsonl").write_text("changed", encoding="utf-8")
        return source

    monkeypatch.setattr(engine, "read_transactions", change_after_read)
    with pytest.raises(BundleError, match="changed"):
        analyze_transactions(fixture.root / "master.csv", bundle=fixture.root)


def test_upstream_workbook_corruption_fails_readback(bundle_factory, tmp_path, monkeypatch):
    fixture = bundle_factory()
    real_write = publication.write_workbook

    def corrupt(plan, staging):
        real_write(plan, staging)
        book = load_workbook(staging)
        sheet = book["Data Quality"]
        first_row = range_boundaries(sheet.tables["UpstreamQuality"].ref)[1]
        sheet.cell(first_row + 1, 2).value = 999
        book.save(staging)
        book.close()

    monkeypatch.setattr(publication, "write_workbook", corrupt)
    with pytest.raises(ReportError, match="verification failed"):
        generate_report(fixture.root / "master.csv", tmp_path / "corrupt.xlsx", bundle=fixture.root)
    assert not (tmp_path / "corrupt.xlsx").exists()


@pytest.mark.parametrize(
    "field,value",
    [("raw_values", "[]"), ("reason_codes", "[]"), ("reasons", "[]"), ("source_row", "2.0")],
)
def test_exclusion_nested_types_and_row_syntax_are_strict(bundle_factory, field, value):
    fixture = bundle_factory()
    rows = fixture.exclusions()
    rows[0][field] = value
    fixture.write_exclusions(rows)
    with pytest.raises(BundleError):
        analyze_transactions(fixture.root / "master.csv", bundle=fixture.root)
