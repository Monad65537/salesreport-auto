"""Focused release regressions for boundary values and publication truthfulness."""

import json
import socket
from decimal import localcontext
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries

from salesreport_auto import (
    InputError,
    ReportError,
    generate_report,
    publication,
    read_transactions,
)
from salesreport_auto.cli import main


def _table(book, sheet_name, table_name):
    sheet = book[sheet_name]
    left, top, right, bottom = range_boundaries(sheet.tables[table_name].ref)
    rows = list(sheet.iter_rows(min_row=top, max_row=bottom, min_col=left, max_col=right))
    return [dict(zip((cell.value for cell in rows[0]), row, strict=True)) for row in rows[1:]]


@pytest.mark.parametrize("mode", ["transactions", "products"])
def test_full_pipeline_is_independent_of_all_decimal_context_traps(
    write_input, write_products, record, product_line, transport, tmp_path, mode
):
    big = "9" * 36 + ".99"
    values = [big, big, "-" + big, "-0.00"]
    rows = []
    for index, value in enumerate(values):
        changes = {"transaction_id": str(index), "amount": value}
        if mode == "products":
            changes.update(
                unit_price=big, quantity="-1" if index == 2 else "0" if index == 3 else "1"
            )
        rows.append((product_line if mode == "products" else record)(**changes))
    source = (write_products if mode == "products" else write_input)(rows, format=transport)
    with localcontext() as context:
        context.prec = 2
        context.Emax = 2
        context.Emin = -2
        for signal in context.traps:
            context.traps[signal] = True
        result = generate_report(source, tmp_path / "exact.xlsx", mode=mode)
        assert not any(context.flags.values())
    metrics = result.analysis.aggregates.summary[0].metrics
    assert metrics.net_revenue_minor == 10**38 - 1
    assert metrics.transaction_count == 4
    assert metrics.average_transaction_value == "25" + "0" * 34 + ".00"
    assert result.review_reasons == ("CHART_PRECISION_LIMIT",)
    book = load_workbook(result.output)
    try:
        assert _table(book, "Summary", "SummaryMetrics")[0]["net_revenue"].value == big
        assert not book["Monthly Sales"]._charts
        assert not book["Branches"]._charts
    finally:
        book.close()


@pytest.mark.parametrize("column", [7, 10, 16384])
@pytest.mark.parametrize("storage", ["inline", "shared", "indented-inline"])
def test_empty_text_beyond_xlsx_headers_is_not_nonempty_data(
    write_input, record, tmp_path, column, storage
):
    source = write_input([record()], format="xlsx")
    book = load_workbook(source)
    book.active.cell(2, column, "").data_type = "s"
    book.save(source)
    book.close()
    if storage != "inline":
        namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        with ZipFile(source) as archive:
            members = {entry.filename: archive.read(entry) for entry in archive.infolist()}
        worksheet = ET.fromstring(members["xl/worksheets/sheet1.xml"])
        cell = list(worksheet.iter(f"{{{namespace}}}c"))[-1]
        cell.clear()
        cell.attrib.update(r=f"{get_column_letter(column)}2")
        if storage == "shared":
            cell.attrib["t"] = "s"
            ET.SubElement(cell, f"{{{namespace}}}v").text = "0"
        else:
            cell.attrib["t"] = "inlineStr"
            inline = ET.SubElement(cell, f"{{{namespace}}}is")
            inline.text = "\n  "
            ET.SubElement(inline, f"{{{namespace}}}t").tail = "\n"
        members["xl/worksheets/sheet1.xml"] = ET.tostring(worksheet)
        if storage == "shared":
            members["xl/sharedStrings.xml"] = (
                f'<sst xmlns="{namespace}" count="1" uniqueCount="1"><si><t></t></si></sst>'
            ).encode()
            content_types = ET.fromstring(members["[Content_Types].xml"])
            ET.SubElement(
                content_types,
                "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
                PartName="/xl/sharedStrings.xml",
                ContentType=(
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"
                ),
            )
            members["[Content_Types].xml"] = ET.tostring(content_types)
        with ZipFile(source, "w", ZIP_DEFLATED) as archive:
            for name, content in members.items():
                archive.writestr(name, content)
    result = generate_report(source, tmp_path / "empty-extra.xlsx")
    assert result.analysis.quality.counts.accepted == 1
    assert len(result.analysis.classified.validated.source.headers) == 6


def test_ignored_worksheets_are_disclosed_in_data_quality(write_input, record, tmp_path):
    source = write_input([record()], format="xlsx", extra_sheets=("Z notes", "A notes"))
    result = generate_report(source, tmp_path / "selected.xlsx")
    book = load_workbook(result.output)
    try:
        context = {
            row["field"].value: row["value"].value
            for row in _table(book, "Data Quality", "ImportContext")
        }
        assert context["selected_sheet"] == "Master"
        assert json.loads(context["ignored_sheets"]) == ["A notes", "Z notes"]
    finally:
        book.close()


@pytest.mark.parametrize("column", [7, 10, 16384])
def test_explicit_empty_xlsx_header_is_still_invalid(write_input, record, column):
    source = write_input([record()], format="xlsx")
    book = load_workbook(source)
    book.active.cell(1, column, "").data_type = "s"
    book.save(source)
    book.close()
    with pytest.raises(InputError):
        read_transactions(source)


@pytest.mark.parametrize("case", ["duplicate", "unsafe", "encrypted"])
def test_malformed_xlsx_package_members_fail_before_decoding(write_input, record, case):
    source = write_input([record()], format="xlsx")
    if case == "encrypted":
        data = bytearray(source.read_bytes())
        directory = data.index(b"PK\x01\x02")
        data[directory + 8] |= 1
        source.write_bytes(data)
    else:
        with ZipFile(source, "a") as archive:
            if case == "duplicate":
                with pytest.warns(UserWarning, match="Duplicate name"):
                    archive.writestr("xl/workbook.xml", archive.read("xl/workbook.xml"))
            else:
                archive.writestr("../outside.xml", b"unsafe")
    with pytest.raises(InputError, match="Duplicate|Unsafe|Encrypted"):
        read_transactions(source)


def test_postpublication_cleanup_failure_is_success_with_explicit_warning(
    write_input, record, tmp_path, monkeypatch, capsys
):
    source = write_input([record()])
    target = tmp_path / "published.xlsx"
    original_unlink = Path.unlink

    def deny_owned_temporary(path, *args, **kwargs):
        if path.name.startswith(".published-"):
            raise PermissionError("Synthetic cleanup denial")
        return original_unlink(path, *args, **kwargs)

    with monkeypatch.context() as patch:
        patch.setattr(Path, "unlink", deny_owned_temporary)
        code = main([
            "--mode", "transactions", "--input", str(source), "--output", str(target)
        ])
    captured = capsys.readouterr()
    assert code == 0
    assert json.loads(captured.out)["status"] == "completed"
    assert "published" in captured.err and "temporary" in captured.err
    assert target.is_file()
    book = load_workbook(target)
    book.close()
    owned = list(tmp_path.glob(".published-*.xlsx"))
    assert len(owned) == 1
    for path in owned:
        path.unlink()


def test_prepublication_cleanup_failure_preserves_original_error(
    write_input, record, tmp_path, monkeypatch
):
    source = write_input([record()])
    original_unlink = Path.unlink

    def deny_owned_temporary(path, *args, **kwargs):
        if path.name.startswith(".failed-"):
            raise PermissionError("Synthetic cleanup denial")
        return original_unlink(path, *args, **kwargs)

    def fail_writer(*args, **kwargs):
        raise ReportError("Synthetic writer failure")

    with monkeypatch.context() as patch:
        patch.setattr(Path, "unlink", deny_owned_temporary)
        patch.setattr(publication, "write_workbook", fail_writer)
        with pytest.raises(ReportError, match="Synthetic writer failure") as error:
            generate_report(source, tmp_path / "failed.xlsx")
    assert any("temporary" in note for note in error.value.__notes__)
    assert not (tmp_path / "failed.xlsx").exists()
    for path in tmp_path.glob(".failed-*.xlsx"):
        path.unlink()


@pytest.mark.parametrize("case", ["completed", "rejected", "chart_limit", "bundle"])
def test_exit_json_and_saved_status_agree_without_network(
    write_input, record, bundle_factory, tmp_path, capsys, monkeypatch, case
):
    def deny_network(*args, **kwargs):
        pytest.fail("Reporting must not access the network")

    monkeypatch.setattr(socket, "create_connection", deny_network)
    monkeypatch.setattr(socket.socket, "connect", deny_network)
    args = []
    if case == "bundle":
        bundle = bundle_factory().root
        source = bundle / "master.csv"
        args = ["--bundle", str(bundle)]
    else:
        amount = (
            "bad" if case == "rejected" else "10000000000000.00"
            if case == "chart_limit" else "0.00"
        )
        source = write_input([record(amount=amount)])
    target = tmp_path / "status.xlsx"
    code = main([
        "--mode", "transactions", "--input", str(source), "--output", str(target), *args
    ])
    captured = capsys.readouterr()
    summary = json.loads(captured.out)
    assert captured.err == ""
    assert code == (0 if case == "completed" else 1)
    book = load_workbook(target)
    try:
        context = {
            row["field"].value: row["value"].value
            for row in _table(book, "Summary", "SummaryContext")
        }
        assert context["status"] == summary["status"]
        assert json.loads(context["review_reasons"]) == summary["review_reasons"]
    finally:
        book.close()


def test_readback_rejects_external_workbook_links(write_input, record, tmp_path, monkeypatch):
    from openpyxl.packaging.relationship import Relationship
    from openpyxl.workbook.external_link.external import ExternalBook, ExternalLink

    source = write_input([record()])
    original_writer = publication.write_workbook

    def inject_link(plan, path):
        original_writer(plan, path)
        book = load_workbook(path)
        link = ExternalLink(externalBook=ExternalBook(id="rId1"))
        link.file_link = Relationship(
            Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/externalLinkPath",
            Target="https://example.invalid/not-opened.xlsx",
            TargetMode="External",
        )
        book._external_links.append(link)
        book.save(path)
        book.close()

    monkeypatch.setattr(publication, "write_workbook", inject_link)
    with pytest.raises(ReportError, match="external workbook links"):
        generate_report(source, tmp_path / "linked-report.xlsx")
    assert not (tmp_path / "linked-report.xlsx").exists()


@pytest.mark.parametrize("mode", ["transactions", "products"])
def test_repeated_publications_have_identical_logical_tables_and_chart_plans(
    write_input, write_products, record, product_line, tmp_path, mode
):
    make = product_line if mode == "products" else record
    rows = [
        make(transaction_id="B", branch_id="002", amount="6.00"),
        make(transaction_id="A", branch_id="001", amount="6.00"),
        make(transaction_id="JPY", currency="JPY", amount="6", unit_price="3"),
    ]
    writer = write_products if mode == "products" else write_input
    source = writer(rows)
    first = generate_report(source, tmp_path / "first.xlsx", mode=mode)
    second = generate_report(source, tmp_path / "second.xlsx", mode=mode)
    assert first.plan == second.plan
    first_book = load_workbook(first.output)
    second_book = load_workbook(second.output)
    try:
        assert first_book.sheetnames == second_book.sheetnames
        for left, right in zip(first_book, second_book, strict=True):
            assert list(left.values) == list(right.values)
            assert len(left._charts) == len(right._charts)
    finally:
        first_book.close()
        second_book.close()
