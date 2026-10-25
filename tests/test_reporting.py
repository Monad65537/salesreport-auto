import json
import os
from dataclasses import replace
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from salesreport_auto import (
    BUSINESS_COLUMNS,
    CurrencyConfig,
    IntegrityError,
    ReportError,
    analyze_transactions,
    build_report_plan,
    generate_report,
    publication,
    publish_report,
    verify_workbook,
    workbook,
)
from salesreport_auto.report_model import CHART_INTEGER_LIMIT, SHEET_NAMES


def _table_rows(book, sheet_name, table_name):
    sheet = book[sheet_name]
    table = sheet.tables[table_name]
    first_col, first_row, last_col, last_row = range_boundaries(table.ref)
    records = list(
        sheet.iter_rows(
            min_col=first_col,
            max_col=last_col,
            min_row=first_row,
            max_row=last_row,
            values_only=True,
        )
    )
    return [dict(zip(records[0], row, strict=True)) for row in records[1:]]


def _title_text(title):
    return "".join(
        node.text or "" for node in title.to_tree().iter() if node.tag.rsplit("}", 1)[-1] == "t"
    )


def test_report_tables_metrics_quality_and_literal_types(write_input, record, transport, tmp_path):
    rows = [
        record(branch_id="A", transaction_id="T1", customer_id="C1", amount="10.00"),
        record(branch_id="A", transaction_id="T2", customer_id="C1", amount="-2.00"),
        record(branch_id="A", transaction_id="T3", customer_id="C2", amount="0.00"),
        record(
            branch_id="B",
            transaction_id="T1",
            customer_id="C1",
            amount="3.01",
            transaction_date="2024-02-01",
        ),
    ]
    source = write_input(rows, format=transport)
    before = source.read_bytes(), source.stat().st_mtime_ns
    result = generate_report(source, tmp_path / "report.xlsx")
    assert result.status == "completed" and result.review_reasons == ()
    book = load_workbook(result.output, data_only=False)
    try:
        assert tuple(book.sheetnames) == SHEET_NAMES
        assert all(sheet.sheet_state == "visible" for sheet in book)
        metrics = _table_rows(book, "Summary", "SummaryMetrics")
        assert metrics == [
            {
                "currency": "USD",
                "precision": 2,
                "accepted_record_count": 4,
                "transaction_count": 4,
                "unique_customer_count": 3,
                "net_revenue": "11.01",
                "average_transaction_value": "2.75",
            }
        ]
        monthly = _table_rows(book, "Monthly Sales", "MonthlyMetrics")
        assert [
            (row["month"], row["net_revenue"], row["average_transaction_value"]) for row in monthly
        ] == [("2024-01", "8.00", "2.67"), ("2024-02", "3.01", "3.01")]
        context = {
            row["field"]: row["value"] for row in _table_rows(book, "Summary", "SummaryContext")
        }
        assert context["upstream_input_count"] == "Unknown / not supplied"
        assert context["local_monetary_complete"] is True
        assert context["input_sha256"] == result.analysis.classified.validated.source.sha256
        assert "Precise totals do not imply" in book["Summary"]["A2"].value
        raw = _table_rows(book, "Raw Data", "ImportedRecords")
        assert len(raw) == 4 and raw[2]["amount"] == "0.00"
        assert raw[1]["report_amount_minor"] == "-200"
        assert _table_rows(book, "Data Quality", "LocalQuality")[0]["accepted"] == 4
        for sheet in book:
            assert sheet.freeze_panes and sheet.auto_filter.ref
            for row in sheet:
                assert all(cell.data_type != "f" for cell in row)
        assert book["Summary"]["F7"].data_type == "s"
        assert book["Summary"]["G7"].data_type == "s"
    finally:
        book.close()
    assert (source.read_bytes(), source.stat().st_mtime_ns) == before
    assert not list(tmp_path.glob(".report-*.xlsx"))
    verify_workbook(result.output, result.plan)


@pytest.mark.parametrize(
    "token", ["=1+1", "+0002", "-0003", "@literal", "000001", "NA", "Łódź\r\n下一行"]
)
def test_formula_like_identifiers_and_raw_text_remain_exact(write_input, record, tmp_path, token):
    source = write_input([record(transaction_id=token, customer_id=token)])
    result = generate_report(source, tmp_path / "literal.xlsx")
    book = load_workbook(result.output, data_only=False)
    try:
        for coordinate in ("B7", "D7"):
            cell = book["Raw Data"][coordinate]
            assert cell.value == token and cell.data_type == "s"
        assert book["Raw Data"]["E7"].value == "12.34"
        assert book["Raw Data"]["E7"].data_type == "s"
        kinds = _table_rows(book, "Raw Data", "ImportedRecords")[0]["report_raw_cell_kinds"]
        assert json.loads(kinds)["transaction_id"] == "text"
    finally:
        book.close()


def test_mixed_currency_charts_have_exact_isolated_sources(write_input, record, tmp_path):
    source = write_input(
        [
            record(transaction_id="u1", amount="-0.01"),
            record(transaction_id="u2", amount="0.00", transaction_date="2024-02-01"),
            record(transaction_id="j1", amount="123", currency="JPY"),
            record(transaction_id="p1", amount="3.25", currency="PLN", branch_id="002"),
        ]
    )
    result = generate_report(source, tmp_path / "mixed.xlsx")
    book = load_workbook(result.output)
    try:
        for sheet_name, kind in (("Monthly Sales", "monthly"), ("Branches", "branches")):
            sheet = book[sheet_name]
            assert len(sheet._charts) == 3
            for chart, currency in zip(sheet._charts, ("JPY", "PLN", "USD"), strict=True):
                assert currency in _title_text(chart.title)
                assert f"{currency} minor units" in _title_text(chart.y_axis.title)
                assert len(chart.series) == 1
                series = chart.series[0]
                assert series.tx.v.startswith(currency)
                assert series.cat.strRef.f.startswith(f"'{sheet_name}'!")
                assert series.val.numRef.f.startswith(f"'{sheet_name}'!")
                first_col, first_row, last_col, last_row = range_boundaries(
                    series.val.numRef.f.split("!")[1]
                )
                assert first_col == last_col == 12
                for number in range(first_row, last_row + 1):
                    assert sheet.cell(number, 10).value == currency
                    assert type(sheet.cell(number, 12).value) is int
                if kind == "branches":
                    assert chart.type == "col"
    finally:
        book.close()


@pytest.mark.parametrize(
    "minor", [CHART_INTEGER_LIMIT, CHART_INTEGER_LIMIT + 1, -(CHART_INTEGER_LIMIT + 1)]
)
def test_chart_precision_boundary_changes_report_only(write_input, record, tmp_path, minor):
    source = write_input([record(currency="JPY", amount=str(minor))])
    analysis = analyze_transactions(source)
    result = publish_report(analysis, tmp_path / "big.xlsx")
    omitted = abs(minor) > CHART_INTEGER_LIMIT
    assert analysis.quality.status == "completed"
    assert result.status == ("review_required" if omitted else "completed")
    assert result.review_reasons == (("CHART_PRECISION_LIMIT",) if omitted else ())
    book = load_workbook(result.output)
    try:
        assert _table_rows(book, "Summary", "SummaryMetrics")[0]["net_revenue"] == str(minor)
        assert len(book["Monthly Sales"]._charts) == int(not omitted)
        assert len(book["Branches"]._charts) == int(not omitted)
        assert book["Monthly Sales"]["L7"].data_type == ("s" if omitted else "n")
        if omitted:
            context = {
                r["field"]: r["value"] for r in _table_rows(book, "Summary", "SummaryContext")
            }
            assert context["status"] == "review_required"
            assert json.loads(context["review_reasons"]) == ["CHART_PRECISION_LIMIT"]
    finally:
        book.close()


def test_record_level_quality_rejections_duplicates_and_warnings(write_input, record, tmp_path):
    headers = [*BUSINESS_COLUMNS, "source_file", "source_sheet", "source_row"]
    source = write_input(
        [
            record(source_file="original.csv", source_row="2"),
            record(source_file="original.csv", source_row="3"),
            record(
                transaction_id="bad-date", transaction_date="bad", amount="3.00", source_row="0"
            ),
            record(transaction_id="bad-money", amount="NaN"),
            record(transaction_id="warning", source_file="only-file.csv"),
        ],
        headers=headers,
    )
    result = generate_report(source, tmp_path / "quality.xlsx")
    assert result.to_summary()["local_counts"] == {
        "input": 5,
        "accepted": 2,
        "rejected": 2,
        "redundant": 1,
    }
    assert result.review_reasons == (
        "LOCAL_REJECTIONS",
        "LOCAL_UNQUANTIFIED_AMOUNTS",
        "INVALID_PROVENANCE",
        "INCOMPLETE_PROVENANCE",
    )
    book = load_workbook(result.output)
    try:
        details = _table_rows(book, "Data Quality", "QualityDetails")
        assert len(details) == 4
        redundant = next(row for row in details if row["status"] == "redundant")
        assert json.loads(redundant["duplicate_of"])["import_row"] == 2
        assert json.loads(redundant["reason_codes"]) == ["EXACT_DUPLICATE"]
        invalid_date = next(row for row in details if "INVALID_DATE" in row["reason_codes"])
        assert json.loads(invalid_date["raw_values"])["transaction_date"] == "bad"
        ledgers = _table_rows(book, "Data Quality", "LocalLedgers")
        usd = next(row for row in ledgers if row["currency"] == "USD")
        assert usd["input_known_amount"] == "40.02"
        assert usd["accepted_known_amount"] == "24.68"
        assert usd["rejected_known_amount"] == "3.00"
        assert usd["redundant_known_amount"] == "12.34"
        assert usd["input_unquantified_count"] == 1
        assert (
            _table_rows(book, "Data Quality", "UpstreamQuality")[0]["rejected"]
            == "Unknown / not supplied"
        )
        raw = _table_rows(book, "Raw Data", "ImportedRecords")
        assert len(raw) == 5 and raw[2]["report_amount_minor"] == "300"
        assert raw[3]["report_amount_minor"] is None
    finally:
        book.close()


@pytest.mark.parametrize("all_rejected", [False, True])
def test_empty_accepted_results_keep_headers_without_charts(
    write_input, record, tmp_path, all_rejected
):
    source = write_input([record(amount="bad")] if all_rejected else [])
    result = generate_report(source, tmp_path / "empty.xlsx")
    assert result.status == "review_required"
    assert result.review_reasons[0] == "NO_ACCEPTED_RECORDS"
    assert result.to_summary()["currencies"] == []
    book = load_workbook(result.output)
    try:
        assert book["Summary"]["A6"].value == "currency"
        assert book["Monthly Sales"]["A6"].value == "currency"
        assert book["Branches"]["A6"].value == "currency"
        assert not book["Monthly Sales"]._charts and not book["Branches"]._charts
        assert "SummaryMetrics" not in book["Summary"].tables
        assert tuple(book.sheetnames) == SHEET_NAMES
    finally:
        book.close()


def test_rounded_negative_ties_and_custom_precision_display(write_input, record, tmp_path):
    source = write_input(
        [
            record(transaction_id="1", amount="0.00", currency="EUR"),
            record(transaction_id="2", amount="-0.01", currency="EUR"),
        ]
    )
    result = generate_report(source, tmp_path / "average.xlsx", config=CurrencyConfig({"EUR": 2}))
    book = load_workbook(result.output)
    try:
        assert (
            _table_rows(book, "Summary", "SummaryMetrics")[0]["average_transaction_value"]
            == "-0.01"
        )
        precision = {
            row["currency"]: row for row in _table_rows(book, "Summary", "SummaryPrecision")
        }
        assert precision["EUR"]["origin"] == "user"
    finally:
        book.close()


def test_existing_report_and_input_alias_are_never_changed(write_input, record, tmp_path):
    source = write_input([record()], format="xlsx")
    original = source.read_bytes()
    with pytest.raises(ReportError, match="alias"):
        generate_report(source, source)
    target = tmp_path / "existing.xlsx"
    target.write_bytes(b"existing report")
    with pytest.raises(ReportError, match="already exists"):
        generate_report(source, target)
    alias = tmp_path / "hardlink.xlsx"
    os.link(source, alias)
    with pytest.raises(ReportError, match="alias"):
        generate_report(source, alias)
    assert target.read_bytes() == b"existing report"
    assert source.read_bytes() == alias.read_bytes() == original


@pytest.mark.parametrize("output", ["report.csv", "missing/report.xlsx"])
def test_output_extension_and_parent_validation(write_input, record, tmp_path, output):
    source = write_input([record()])
    with pytest.raises(ReportError):
        generate_report(source, tmp_path / output)
    assert not (tmp_path / output).exists()


def test_writer_failure_cleans_only_owned_staging_file(write_input, record, tmp_path, monkeypatch):
    source = write_input([record()])
    sentinel = tmp_path / ".report-someone-else.xlsx"
    sentinel.write_bytes(b"unrelated temporary file")

    def fail(plan, path):
        Path(path).write_bytes(b"partial")
        raise OSError("synthetic disk failure")

    monkeypatch.setattr(publication, "write_workbook", fail)
    with pytest.raises(ReportError):
        generate_report(source, tmp_path / "report.xlsx")
    assert not (tmp_path / "report.xlsx").exists()
    assert list(tmp_path.glob(".report-*.xlsx")) == [sentinel]
    assert sentinel.read_bytes() == b"unrelated temporary file"


def test_atomic_publication_collision_preserves_competing_report(
    write_input, record, tmp_path, monkeypatch
):
    source = write_input([record()])
    target = tmp_path / "race.xlsx"
    real_publish = publication._publish_no_replace

    def collision(staging, destination):
        destination.write_bytes(b"another publisher")
        real_publish(staging, destination)

    monkeypatch.setattr(publication, "_publish_no_replace", collision)
    with pytest.raises(ReportError):
        generate_report(source, target)
    assert target.read_bytes() == b"another publisher"
    assert not list(tmp_path.glob(".race-*.xlsx"))


@pytest.mark.parametrize("failure", ["value", "formula", "chart", "sheet", "quality"])
def test_readback_corruption_prevents_publication(
    write_input, record, tmp_path, monkeypatch, failure
):
    source = write_input([record()])
    real_write = publication.write_workbook

    def corrupt(plan, staging):
        real_write(plan, staging)
        book = load_workbook(staging)
        if failure == "value":
            book["Summary"]["F7"] = "999.00"
        elif failure == "formula":
            book["Raw Data"]["A7"] = "=1+1"
        elif failure == "chart":
            book["Monthly Sales"]._charts[0].series[0].val.numRef.f = "'Branches'!$L$7"
        elif failure == "sheet":
            book.create_sheet("Products")
        else:
            book["Data Quality"]["C7"] = 999
        book.save(staging)
        book.close()

    monkeypatch.setattr(publication, "write_workbook", corrupt)
    with pytest.raises(ReportError):
        generate_report(source, tmp_path / "broken.xlsx")
    assert not (tmp_path / "broken.xlsx").exists()
    assert not list(tmp_path.glob(".broken-*.xlsx"))


def test_input_change_during_render_prevents_publication(
    write_input, record, tmp_path, monkeypatch
):
    source = write_input([record()])
    real_write = publication.write_workbook

    def mutate_source(plan, staging):
        real_write(plan, staging)
        source.write_bytes(source.read_bytes() + b"\n")

    monkeypatch.setattr(publication, "write_workbook", mutate_source)
    with pytest.raises(ReportError, match="Input changed"):
        generate_report(source, tmp_path / "changed.xlsx")
    assert not (tmp_path / "changed.xlsx").exists()
    assert not list(tmp_path.glob(".changed-*.xlsx"))


@pytest.mark.parametrize(
    "value",
    ["a" * 32768, "😀" * 16384, "illegal\x01", "illegal\uffff"],
    ids=("overlong_ascii", "overlong_utf16", "control_character", "xml_noncharacter"),
)
def test_excel_text_limits_fail_without_truncating(write_input, record, tmp_path, value):
    source = write_input([record(customer_id=value)])
    with pytest.raises(ReportError, match="Excel"):
        generate_report(source, tmp_path / "limit.xlsx")
    assert not (tmp_path / "limit.xlsx").exists()
    assert not list(tmp_path.glob(".limit-*.xlsx"))


def test_excel_row_limit_and_stale_analysis_fail_before_publication(
    write_input, record, tmp_path, monkeypatch
):
    source = write_input([record()])
    analysis = analyze_transactions(source)
    broken = replace(analysis, aggregates=replace(analysis.aggregates, summary=()))
    with pytest.raises(IntegrityError):
        build_report_plan(broken, application_version="test")
    monkeypatch.setattr(workbook, "EXCEL_MAX_ROWS", 10)
    with pytest.raises(ReportError, match="row limit"):
        publish_report(analysis, tmp_path / "rows.xlsx")
    assert not (tmp_path / "rows.xlsx").exists()


def test_report_zip_has_no_formula_cells_and_exact_helper_xml(write_input, record, tmp_path):
    source = write_input([record(currency="JPY", amount=str(CHART_INTEGER_LIMIT))])
    result = generate_report(source, tmp_path / "exact.xlsx")
    with ZipFile(result.output) as archive:
        xml = archive.read("xl/worksheets/sheet2.xml")
        assert b"<" in xml and str(CHART_INTEGER_LIMIT).encode() in xml
    verify_workbook(result.output, result.plan)


def test_precision_limit_is_applied_per_chart_not_per_currency(write_input, record, tmp_path):
    source = write_input(
        [
            record(branch_id="A", transaction_id="1", currency="JPY", amount="500000000000000"),
            record(branch_id="B", transaction_id="2", currency="JPY", amount="500000000000000"),
        ]
    )
    result = generate_report(source, tmp_path / "partial-charts.xlsx")
    assert result.plan.omitted_charts == ("Monthly Sales: JPY",)
    assert result.review_reasons == ("CHART_PRECISION_LIMIT",)
    book = load_workbook(result.output)
    try:
        assert not book["Monthly Sales"]._charts
        assert len(book["Branches"]._charts) == 1
        assert (
            _table_rows(book, "Summary", "SummaryMetrics")[0]["net_revenue"] == "1000000000000000"
        )
    finally:
        book.close()


def test_unsupported_hardlink_filesystem_fails_closed(write_input, record, tmp_path, monkeypatch):
    source = write_input([record()])

    def unsupported(staging, target):
        raise OSError("Hard links are unsupported")

    monkeypatch.setattr(publication, "_publish_no_replace", unsupported)
    with pytest.raises(ReportError):
        generate_report(source, tmp_path / "unsupported.xlsx")
    assert not (tmp_path / "unsupported.xlsx").exists()
    assert not list(tmp_path.glob(".unsupported-*.xlsx"))


def test_output_link_paths_are_rejected_without_real_symlink_privileges(
    write_input, record, tmp_path, monkeypatch
):
    source = write_input([record()])
    target = tmp_path / "dangling.xlsx"
    real_is_symlink = Path.is_symlink
    monkeypatch.setattr(Path, "is_symlink", lambda path: path == target or real_is_symlink(path))
    with pytest.raises(ReportError, match="symbolic link"):
        generate_report(source, target)
    assert not target.exists()


def test_native_rejected_xlsx_values_are_textual_audit_data(write_input, record, tmp_path):
    source = write_input([record()], format="xlsx")
    book = load_workbook(source)
    book["Master"]["E2"] = 1.25
    book.save(source)
    book.close()
    result = generate_report(source, tmp_path / "native-rejected.xlsx")
    assert result.status == "review_required"
    book = load_workbook(result.output)
    try:
        assert book["Raw Data"]["E7"].value == "1.25"
        assert book["Raw Data"]["E7"].data_type == "s"
        raw = _table_rows(book, "Raw Data", "ImportedRecords")[0]
        assert raw["report_amount_minor"] is None
        assert json.loads(raw["report_raw_cell_kinds"])["amount"] == "numeric"
    finally:
        book.close()
