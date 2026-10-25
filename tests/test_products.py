from dataclasses import replace
from decimal import Decimal, localcontext

import pytest
from openpyxl import load_workbook

from salesreport_auto import (
    PRODUCT_ALL_COLUMNS,
    PRODUCT_COLUMNS,
    ConfigurationError,
    CurrencyConfig,
    InputError,
    IntegrityError,
    OutcomeAmounts,
    OutcomeCounts,
    analyze_products,
    analyze_transactions,
    classify_products,
    product_sales,
    read_products,
    read_transactions,
    reconcile_transactions,
    validate_products,
    validate_transactions,
)


def test_s16_product_lines_and_distinct_transaction_denominators(
    write_products, product_line, transport
):
    source = write_products(
        [
            product_line(transaction_id="T1"),
            product_line(
                transaction_id="T1",
                line_id="L2",
                product_id="P2",
                quantity="1",
                unit_price="4.00",
                amount="4.00",
            ),
            product_line(transaction_id="T2", quantity="-1", amount="-3.00"),
        ],
        format=transport,
    )
    analysis = analyze_products(source)
    assert analysis.mode == "products" and analysis.upstream is None
    assert analysis.quality.counts == OutcomeCounts(3, 3, 0, 0)
    assert not analysis.quality.review_required
    summary = analysis.aggregates.summary[0].metrics
    assert (
        summary.accepted_record_count,
        summary.transaction_count,
        summary.unique_customer_count,
    ) == (3, 2, 1)
    assert (summary.net_revenue, summary.average_transaction_value) == ("7.00", "3.50")
    assert (summary.average.numerator_minor, summary.average.denominator) == (700, 2)
    assert analysis.aggregates.monthly[0].metrics == summary
    assert analysis.aggregates.branches[0].metrics == summary
    second, first = analysis.aggregates.products
    assert (second.product_id, second.revenue_rank, second.net_revenue) == ("P2", 1, "4.00")
    assert (first.product_id, first.revenue_rank, first.net_revenue) == ("P1", 2, "3.00")
    assert (first.positive_quantity, first.negative_quantity_magnitude, first.net_quantity) == (
        2,
        1,
        1,
    )
    assert first.accepted_line_count == first.transaction_count == 2
    assert first.unique_customer_count == 1
    assert first.minimum_unit_price == first.maximum_unit_price == "3.00"
    assert first.positive_quantity_weighted_unit_price == "3.00"
    assert sum(row.transaction_count for row in analysis.aggregates.products) == 3
    assert product_sales(analysis.classified) == analysis.aggregates.products


def test_s17_duplicates_are_line_scoped_and_ignore_provenance(
    write_products, product_line, transport
):
    headers = [*PRODUCT_COLUMNS, "source_file", "source_row"]
    source = write_products(
        [
            product_line(source_file="z.csv", source_row="90"),
            product_line(
                quantity="+0002",
                amount="6.000",
                unit_price="+3.000",
                source_file="a.csv",
                source_row="2",
            ),
            product_line(line_id="L2", source_file="a.csv", source_row="3"),
            product_line(transaction_id="another", source_file="a.csv", source_row="4"),
        ],
        headers=headers,
        format=transport,
    )
    analysis = analyze_products(source)
    assert analysis.quality.counts == OutcomeCounts(4, 3, 0, 1)
    assert not analysis.quality.review_required
    assert analysis.classified.redundant[0].duplicate_of.import_row == 2
    assert analysis.classified.accepted[0].row.provenance.source_file == "z.csv"
    metrics = analysis.aggregates.summary[0].metrics
    assert (metrics.transaction_count, metrics.net_revenue, metrics.average_transaction_value) == (
        2,
        "18.00",
        "9.00",
    )
    assert analysis.aggregates.products[0].accepted_line_count == 3


@pytest.mark.parametrize(
    "change,code,unquantified",
    [
        ({"quantity": "1.5"}, "INVALID_QUANTITY", 0),
        ({"quantity": "1000000001"}, "INVALID_QUANTITY", 0),
        ({"quantity": "1e2"}, "INVALID_QUANTITY", 0),
        ({"quantity": " 2"}, "INVALID_QUANTITY", 0),
        ({"quantity": ""}, "MISSING_REQUIRED_VALUE", 0),
        ({"unit_price": "-3.00"}, "NEGATIVE_UNIT_PRICE", 0),
        ({"unit_price": "0.001"}, "EXCESS_FRACTIONAL_PRECISION", 0),
        ({"unit_price": "NaN"}, "INVALID_AMOUNT", 0),
        ({"unit_price": "1e2"}, "INVALID_AMOUNT", 0),
        ({"unit_price": ""}, "MISSING_REQUIRED_VALUE", 0),
        ({"line_id": ""}, "MISSING_REQUIRED_VALUE", 0),
        ({"product_id": ""}, "MISSING_REQUIRED_VALUE", 0),
        ({"transaction_date": "bad"}, "INVALID_DATE", 0),
        ({"customer_id": ""}, "MISSING_REQUIRED_VALUE", 0),
        ({"amount": "bad"}, "INVALID_AMOUNT", 1),
        ({"currency": "EUR"}, "UNSUPPORTED_CURRENCY", 1),
        ({"amount": "6.01"}, "LINE_AMOUNT_MISMATCH", 0),
    ],
)
def test_s17_one_invalid_line_quarantines_complete_transaction(
    write_products, product_line, transport, change, code, unquantified
):
    source = write_products(
        [
            product_line(),
            product_line(),
            product_line(**{"line_id": "L2", **change}),
            product_line(transaction_id="safe"),
        ],
        format=transport,
    )
    analysis = analyze_products(source)
    assert analysis.quality.counts == OutcomeCounts(4, 1, 3, 0)
    assert analysis.quality.unquantified_count == unquantified
    assert all(
        "TRANSACTION_HAS_INVALID_LINE" in {i.code for i in item.issues}
        for item in analysis.classified.rejected
    )
    assert code in {issue.code for issue in analysis.classified.rejected[2].issues}
    assert analysis.classified.accepted[0].row.transaction_id == "safe"
    assert analysis.aggregates.summary[0].metrics.net_revenue == "6.00"
    assert analysis.quality.count_balanced and analysis.quality.known_amounts_balanced


@pytest.mark.parametrize(
    "change",
    [
        {"quantity": "1", "amount": "3.00"},
        {"product_id": "other"},
        {"product_name": "Label"},
    ],
)
def test_s17_conflicting_line_rejects_order_including_equal_copies(
    write_products, product_line, transport, change
):
    source = write_products(
        [
            product_line(),
            product_line(),
            product_line(**change),
            product_line(line_id="L2"),
        ],
        headers=[*PRODUCT_COLUMNS, "product_name"],
        format=transport,
    )
    analysis = analyze_products(source)
    assert analysis.quality.counts == OutcomeCounts(4, 0, 4, 0)
    assert all(
        "CONFLICTING_LINE" in {i.code for i in item.issues} for item in analysis.classified.rows
    )
    assert analysis.quality.local_monetary_complete
    assert analysis.aggregates.products == ()


@pytest.mark.parametrize(
    "change",
    [
        {"transaction_date": "2024-02-01"},
        {"customer_id": "other"},
        {"currency": "PLN"},
    ],
)
def test_s17_transaction_metadata_must_agree(write_products, product_line, transport, change):
    analysis = analyze_products(
        write_products(
            [
                product_line(),
                product_line(line_id="L2", **change),
            ],
            format=transport,
        )
    )
    assert analysis.quality.counts == OutcomeCounts(2, 0, 2, 0)
    assert all(
        "INCONSISTENT_TRANSACTION" in {i.code for i in item.issues}
        for item in analysis.classified.rows
    )


def test_unreadable_transaction_key_is_not_guessed(write_products, product_line, transport):
    analysis = analyze_products(
        write_products(
            [
                product_line(),
                product_line(line_id="L2"),
                product_line(branch_id=""),
            ],
            format=transport,
        )
    )
    assert analysis.quality.counts == OutcomeCounts(3, 2, 1, 0)
    assert analysis.classified.rejected[0].row.business_key is None
    assert analysis.aggregates.summary[0].metrics.transaction_count == 1


@pytest.mark.parametrize(
    "quantity,price,amount",
    [
        ("0", "3.00", "0.00"),
        ("-0", "0.00", "-0.00"),
        ("2", "0", "0.00"),
        ("-2", "0", "0.00"),
        ("-2", "3", "-6.00"),
        ("1000000000", "0", "0"),
        ("-1000000000", "0", "0"),
        ("1", "-0.00", "0"),
    ],
)
def test_s18_signed_quantities_and_zero_values_are_valid(
    write_products, product_line, transport, quantity, price, amount
):
    analysis = analyze_products(
        write_products(
            [
                product_line(quantity=quantity, unit_price=price, amount=amount),
            ],
            format=transport,
        )
    )
    assert analysis.quality.counts == OutcomeCounts(1, 1, 0, 0)
    product = analysis.aggregates.products[0]
    assert product.net_quantity == int(quantity)
    assert product.transaction_count == 1
    assert product.positive_quantity_weighted_unit_price == (
        product.minimum_unit_price if int(quantity) > 0 else "N/A"
    )


@pytest.mark.parametrize("quantity,amount", [("0", "6.00"), ("-2", "6.00"), ("2", "-6.00")])
def test_s18_quantity_amount_sign_or_zero_mismatches_are_rejected(
    write_products, product_line, transport, quantity, amount
):
    analysis = analyze_products(
        write_products([product_line(quantity=quantity, amount=amount)], format=transport)
    )
    assert analysis.quality.unquantified_count == 0
    assert "LINE_AMOUNT_MISMATCH" in {i.code for i in analysis.classified.rejected[0].issues}


def test_s18_weighted_price_uses_positive_units_only(write_products, product_line, transport):
    analysis = analyze_products(
        write_products(
            [
                product_line(transaction_id="A", quantity="1", unit_price="0.01", amount="0.01"),
                product_line(transaction_id="B", quantity="3", unit_price="0.02", amount="0.06"),
                product_line(transaction_id="C", quantity="-2", unit_price="0.99", amount="-1.98"),
                product_line(transaction_id="D", quantity="0", unit_price="2.00", amount="0.00"),
            ],
            format=transport,
        )
    )
    product = analysis.aggregates.products[0]
    assert (
        product.positive_quantity,
        product.negative_quantity_magnitude,
        product.net_quantity,
    ) == (4, 2, 2)
    assert product.net_revenue == "-1.91"
    assert (product.minimum_unit_price, product.maximum_unit_price) == ("0.01", "2.00")
    assert product.positive_quantity_weighted_unit_price == "0.02"
    assert (
        product.positive_quantity_average.numerator_minor,
        product.positive_quantity_average.denominator,
    ) == (7, 4)


def test_s19_product_scope_ranking_and_label_conflicts(write_products, product_line, transport):
    rows = [
        product_line(transaction_id="T1", product_id="2", product_name="Zulu"),
        product_line(transaction_id="T2", branch_id="002", product_id="2", product_name="Alpha"),
        product_line(transaction_id="T3", product_id="10", quantity="4", amount="12.00"),
        product_line(
            transaction_id="T4", product_id="2", currency="PLN", product_name="Different currency"
        ),
        product_line(transaction_id="T5", product_id="zero", quantity="0", amount="0.00"),
        product_line(transaction_id="T6", product_id="negative", quantity="-1", amount="-3.00"),
    ]
    analysis = analyze_products(
        write_products(rows, headers=[*PRODUCT_COLUMNS, "product_name"], format=transport)
    )
    assert analysis.quality.review_reasons == ("PRODUCT_LABEL_CONFLICT",)
    products = analysis.aggregates.products
    assert [(p.currency, p.product_id, p.revenue_rank) for p in products] == [
        ("PLN", "2", 1),
        ("USD", "10", 1),
        ("USD", "2", 2),
        ("USD", "zero", 3),
        ("USD", "negative", 4),
    ]
    product = next(p for p in products if p.currency == "USD" and p.product_id == "2")
    assert product.product_name == "Alpha" and product.unique_customer_count == 2
    assert product.transaction_count == 2 and product.net_revenue == "12.00"
    assert next(p for p in products if p.product_id == "10").product_name == "10"
    warned = [item for item in analysis.classified.rows if item.row.warnings]
    assert len(warned) == 2
    assert classify_products(analysis.classified.validated) == analysis.classified


def test_optional_absent_label_does_not_conflict_with_a_supplied_label(
    write_products, product_line, transport
):
    analysis = analyze_products(
        write_products(
            [
                product_line(product_name="   "),
                product_line(line_id="L2", product_name="Named"),
            ],
            headers=[*PRODUCT_COLUMNS, "product_name"],
            format=transport,
        )
    )
    assert not analysis.quality.review_required
    assert analysis.aggregates.products[0].product_name == "Named"


@pytest.mark.parametrize("label,valid", [("字" * 200, True), ("字" * 201, False)])
def test_optional_product_label_length(write_products, product_line, transport, label, valid):
    source = write_products(
        [product_line(product_name=label)],
        headers=[*PRODUCT_COLUMNS, "product_name"],
        format=transport,
    )
    analysis = analyze_products(source)
    assert analysis.quality.counts.accepted == int(valid)
    if not valid:
        assert "INVALID_PRODUCT_NAME" in {i.code for i in analysis.classified.rejected[0].issues}


def test_s20_mode_schemas_are_explicit_not_inferred(
    write_input, write_products, product_line, record, transport
):
    transactions = write_input([record()], format=transport, name=f"transactions.{transport}")
    products = write_products([product_line()], format=transport, name=f"products.{transport}")
    with pytest.raises(InputError, match="missing=.*line_id"):
        analyze_products(transactions)
    with pytest.raises(InputError, match="unknown=.*line_id|beyond supported headers"):
        analyze_transactions(products)
    with pytest.raises(ConfigurationError):
        validate_transactions(read_products(products))
    with pytest.raises(ConfigurationError):
        validate_products(read_transactions(transactions))


@pytest.mark.parametrize("field", PRODUCT_COLUMNS)
def test_each_required_product_header_is_mandatory(write_products, product_line, field):
    headers = [header for header in PRODUCT_COLUMNS if header != field]
    with pytest.raises(InputError, match="Header mismatch"):
        read_products(write_products([product_line()], headers=headers))


def test_optional_headers_order_unicode_and_raw_values(write_products, product_line, transport):
    row = product_line(
        line_id="+0001",
        product_id="=商品",
        quantity="+0002",
        unit_price="3.000",
        product_name="@Literal",
        source_file="source.csv",
        source_row="0002",
        source_sheet="",
    )
    source = write_products([row], headers=list(reversed(PRODUCT_ALL_COLUMNS)), format=transport)
    analysis = analyze_products(source)
    assert dict(analysis.classified.accepted[0].row.raw.raw_values) == row
    assert analysis.classified.accepted[0].row.line_key == ("001", "000001", "+0001")
    assert analysis.aggregates.products[0].product_id == "=商品"


def test_product_xlsx_requires_one_sheet_even_when_master_exists(write_products, product_line):
    source = write_products([product_line()], format="xlsx", extra_sheets=("Notes",))
    with pytest.raises(InputError, match="exactly one worksheet"):
        read_products(source)


@pytest.mark.parametrize(
    "field,value,code",
    [
        ("line_id", 1, "ID_NOT_TEXT"),
        ("product_id", 2, "ID_NOT_TEXT"),
        ("quantity", 2, "INVALID_CELL_TYPE"),
        ("unit_price", 3.0, "INVALID_CELL_TYPE"),
        ("quantity", True, "INVALID_CELL_TYPE"),
        ("unit_price", "=1+1", "FORMULA_NOT_ALLOWED"),
    ],
)
def test_product_xlsx_business_cells_must_be_text(write_products, product_line, field, value, code):
    source = write_products([product_line()], format="xlsx")
    book = load_workbook(source)
    book["Master"].cell(2, PRODUCT_COLUMNS.index(field) + 1).value = value
    book.save(source)
    book.close()
    analysis = analyze_products(source)
    assert code in {issue.code for issue in analysis.classified.rejected[0].issues}
    assert analysis.quality.unquantified_count == 0


def test_large_money_and_quantity_are_context_independent(write_products, product_line, transport):
    with localcontext() as context:
        context.prec = 2
        analysis = analyze_products(
            write_products(
                [
                    product_line(
                        transaction_id="A",
                        quantity="1",
                        unit_price="9" * 36 + ".99",
                        amount="9" * 36 + ".99",
                    ),
                    product_line(
                        transaction_id="B",
                        quantity="1",
                        unit_price="9" * 36 + ".99",
                        amount="9" * 36 + ".99",
                    ),
                ],
                format=transport,
            )
        )
    product = analysis.aggregates.products[0]
    assert product.net_revenue_minor == 2 * (10**38 - 1)
    assert product.positive_quantity_weighted_unit_price == "9" * 36 + ".99"
    assert analysis.classified.accepted[0].row.unit_price == Decimal("9" * 36 + ".99")


def test_custom_precision_and_currency_separation(write_products, product_line, transport):
    source = write_products(
        [
            product_line(
                transaction_id="A", quantity="3", currency="KWD", unit_price="0.001", amount="0.003"
            ),
            product_line(
                transaction_id="B", quantity="2", currency="JPY", unit_price="3", amount="6"
            ),
        ],
        format=transport,
    )
    analysis = analyze_products(source, config=CurrencyConfig({"KWD": 3}))
    assert [
        (p.currency, p.net_revenue, p.minimum_unit_price) for p in analysis.aggregates.products
    ] == [
        ("JPY", "6", "3"),
        ("KWD", "0.003", "0.001"),
    ]
    assert analysis.quality.known_amounts_balanced


def test_product_conservation_accounts_for_invalid_prices_separately(
    write_products, product_line, transport
):
    analysis = analyze_products(
        write_products(
            [
                product_line(),
                product_line(),
                product_line(transaction_id="bad-price", unit_price="NaN"),
                product_line(transaction_id="bad-money", amount="NaN"),
            ],
            format=transport,
        )
    )
    usd = next(row for row in analysis.quality.currencies if row.currency == "USD")
    assert analysis.quality.counts == OutcomeCounts(4, 1, 2, 1)
    assert analysis.quality.unquantified_count == 1
    assert usd.known_amount_minor == OutcomeAmounts(1800, 600, 600, 600)
    assert usd.unquantified_counts == OutcomeCounts(1, 0, 1, 0)
    assert analysis.aggregates.products[0].net_revenue_minor == usd.known_amount_minor.accepted


def test_product_integrity_checks_detect_partial_acceptance_and_wrong_prices(
    write_products, product_line
):
    source = write_products([product_line(), product_line(line_id="L2", unit_price="NaN")])
    result = analyze_products(source).classified
    broken = replace(
        result, rows=(replace(result.rows[0], status="accepted", issues=()), result.rows[1])
    )
    with pytest.raises(IntegrityError, match="partially quarantined"):
        reconcile_transactions(broken)
    source = write_products([product_line()])
    valid = analyze_products(source).classified
    altered = replace(valid.validated.rows[0], unit_price_minor=301)
    corrupt = replace(
        valid,
        validated=replace(valid.validated, rows=(altered,)),
        rows=(replace(valid.rows[0], row=altered),),
    )
    with pytest.raises(IntegrityError, match="arithmetic"):
        reconcile_transactions(corrupt)


def test_transaction_and_customer_identifiers_are_branch_scoped_for_products(
    write_products, product_line, transport
):
    analysis = analyze_products(
        write_products(
            [
                product_line(branch_id="A", transaction_id="T1", customer_id="C1"),
                product_line(branch_id="B", transaction_id="T1", customer_id="C1"),
            ],
            format=transport,
        )
    )
    metrics = analysis.aggregates.summary[0].metrics
    assert (
        metrics.accepted_record_count,
        metrics.transaction_count,
        metrics.unique_customer_count,
    ) == (2, 2, 2)
    assert analysis.aggregates.products[0].transaction_count == 2
    assert analysis.aggregates.products[0].unique_customer_count == 2
    assert metrics.average_transaction_value == "6.00"


def test_product_representatives_and_rankings_are_deterministic(write_products, product_line):
    rows = [
        product_line(transaction_id="A", product_id="2"),
        product_line(transaction_id="A", product_id="2"),
        product_line(transaction_id="B", product_id="10"),
    ]
    source = write_products(rows)
    first = analyze_products(source)
    assert first == analyze_products(source)
    classified = classify_products(
        replace(first.classified.validated, rows=tuple(reversed(first.classified.validated.rows)))
    )
    assert classified.rows == first.classified.rows
    assert reconcile_transactions(classified) == first.quality
    reordered = analyze_products(write_products(list(reversed(rows))))
    assert reordered.aggregates == first.aggregates
    assert first.classified.redundant[0].duplicate_of.import_row == 2


def test_long_zero_padded_quantities_do_not_add_hidden_length_rules(write_products, product_line):
    source = write_products([product_line(quantity="+" + "0" * 4500 + "2")])
    analysis = analyze_products(source)
    assert analysis.classified.accepted[0].row.quantity == 2
    assert not analysis.quality.review_required


@pytest.mark.parametrize(
    "field,value,code",
    [
        ("unit_price", "9" * 37 + ".99", "AMOUNT_TOO_LARGE"),
        ("unit_price", "0" * 513, "INVALID_AMOUNT"),
        ("quantity", "-1000000001", "INVALID_QUANTITY"),
        ("quantity", "٢", "INVALID_QUANTITY"),
    ],
)
def test_product_price_and_quantity_boundaries(write_products, product_line, field, value, code):
    analysis = analyze_products(write_products([product_line(**{field: value})]))
    assert analysis.quality.counts.accepted == 0 and analysis.quality.unquantified_count == 0
    assert code in {issue.code for issue in analysis.classified.rejected[0].issues}


def test_all_product_group_reasons_and_original_failures_survive(write_products, product_line):
    analysis = analyze_products(
        write_products(
            [
                product_line(),
                product_line(quantity="1", amount="3.00"),
                product_line(line_id="L2", quantity="bad", customer_id="other"),
            ]
        )
    )
    for item in analysis.classified.rejected:
        assert {"TRANSACTION_HAS_INVALID_LINE", "CONFLICTING_LINE", "INCONSISTENT_TRANSACTION"} <= {
            issue.code for issue in item.issues
        }
    assert "INVALID_QUANTITY" in {issue.code for issue in analysis.classified.rejected[2].issues}
    assert analysis.quality.counts == OutcomeCounts(3, 0, 3, 0)


def test_product_header_only_and_blank_occurrences_are_distinct(
    write_products, product_line, transport
):
    empty = analyze_products(write_products([], format=transport))
    assert empty.quality.counts == OutcomeCounts(0, 0, 0, 0)
    assert empty.aggregates.products == ()
    blank = analyze_products(write_products([[""] * 10, product_line()], format=transport))
    assert blank.quality.counts == OutcomeCounts(2, 1, 1, 0)
    assert blank.quality.unquantified_count == 1
    assert {issue.code for issue in blank.classified.rejected[0].issues} == {"BLANK_ROW"}
