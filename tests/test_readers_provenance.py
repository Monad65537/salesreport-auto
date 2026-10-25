import csv
from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZIP_BZIP2, ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import load_workbook
from openpyxl.styles import Font

from salesreport_auto import (
    BUSINESS_COLUMNS,
    InputError,
    InputLimits,
    analyze_transactions,
    read_transactions,
    readers,
)


@pytest.mark.parametrize(
    "headers",
    [
        ["branch_id", "transaction_id"],
        [*BUSINESS_COLUMNS, "product_id"],
        [*BUSINESS_COLUMNS, "amount"],
        [" Branch_id", *BUSINESS_COLUMNS[1:]],
        [*BUSINESS_COLUMNS, ""],
        ["BRANCH_ID", *BUSINESS_COLUMNS[1:]],
    ],
)
def test_s11_bad_headers_fail_structurally(write_input, transport, headers):
    with pytest.raises(InputError, match="Header|headers"):
        read_transactions(write_input([], headers=headers, format=transport))


def test_headers_can_be_reordered_without_changing_raw_order(write_input, record, transport):
    headers = list(reversed(BUSINESS_COLUMNS))
    result = analyze_transactions(write_input([record()], headers=headers, format=transport))
    assert tuple(result.classified.rows[0].row.raw.raw_values) == tuple(headers)
    assert result.aggregates.summary[0].metrics.net_revenue == "12.34"


@pytest.mark.parametrize(
    "contents",
    [
        b"",
        b"\n",
        b"\xff\xfe",
        b"branch_id,transaction_id\n",
        (",".join(BUSINESS_COLUMNS) + '\n"unterminated').encode(),
        (",".join(BUSINESS_COLUMNS) + '\n"quoted"junk,x,x,x,x,x\n').encode(),
        (",".join(BUSINESS_COLUMNS) + "\n1,2,3\n").encode(),
        (",".join(BUSINESS_COLUMNS) + "\n1,2,3,4,5,6,7\n").encode(),
    ],
)
def test_empty_malformed_encoding_and_width_errors(tmp_path, contents):
    path = tmp_path / "bad.csv"
    path.write_bytes(contents)
    with pytest.raises(InputError):
        read_transactions(path)


def test_bom_crlf_unicode_literal_na_and_long_ids(write_input, record):
    row = record(transaction_id="0" * 100, branch_id="Łódź", customer_id="NA")
    path = write_input([row])
    path.write_bytes(b"\xef\xbb\xbf" + path.read_bytes())
    result = analyze_transactions(path)
    assert dict(result.classified.accepted[0].row.raw.raw_values) == row
    assert result.classified.accepted[0].row.customer_id == "NA"


def test_master_selected_exactly_and_other_sheets_recorded(write_input, record):
    path = write_input([record()], format="xlsx", extra_sheets=("Notes", "Archive"))
    raw = read_transactions(path)
    assert raw.selected_sheet == "Master"
    assert raw.ignored_sheets == ("Archive", "Notes")
    assert len(raw.records) == 1
    path = write_input([record()], format="xlsx", sheet="Only")
    assert read_transactions(path).selected_sheet == "Only"
    path = write_input([record()], format="xlsx", sheet="master", extra_sheets=("Notes",))
    with pytest.raises(InputError, match="Master"):
        read_transactions(path)


def test_corrupt_workbook_and_unsupported_file(tmp_path):
    path = tmp_path / "bad.xlsx"
    path.write_bytes(b"not a workbook")
    with pytest.raises(InputError):
        read_transactions(path)
    with pytest.raises(InputError):
        read_transactions(tmp_path / "missing.csv")
    with pytest.raises(InputError):
        read_transactions(tmp_path)
    with pytest.raises(InputError):
        read_transactions(tmp_path / "input.xls")


def test_record_byte_and_archive_limits(write_input, record):
    path = write_input([record(), record(transaction_id="two")])
    with pytest.raises(InputError, match="max_records"):
        read_transactions(path, limits=InputLimits(max_records=1))
    assert len(read_transactions(path, limits=InputLimits(max_records=2)).records) == 2
    length = path.stat().st_size
    assert len(read_transactions(path, limits=InputLimits(max_file_bytes=length)).records) == 2
    with pytest.raises(InputError, match="max_file_bytes"):
        read_transactions(path, limits=InputLimits(max_file_bytes=length - 1))
    book = write_input([record()], format="xlsx")
    with pytest.raises(InputError, match="max_xlsx_uncompressed_bytes"):
        read_transactions(book, limits=InputLimits(max_xlsx_uncompressed_bytes=100))


def test_xlsx_internal_gaps_and_formatted_trailing_rows(write_input, record):
    path = write_input([record(), [""] * 6, record(transaction_id="two")], format="xlsx")
    book = load_workbook(path)
    book["Master"]["A1000000"].font = Font(bold=True)
    book["Master"]["XFD1000000"].font = Font(bold=True)
    book.save(path)
    book.close()
    result = analyze_transactions(path, limits=InputLimits(max_records=3))
    assert [r.row.raw.location.import_row for r in result.classified.rows] == [2, 3, 4]
    assert result.quality.counts.rejected == 1
    assert result.classified.rejected[0].issues[0].code == "BLANK_ROW"


def test_xlsx_large_sparse_span_rejected_before_expansion(write_input, record):
    path = write_input([record()], format="xlsx")
    book = load_workbook(path)
    book["Master"]["A100002"] = "value"
    book.save(path)
    book.close()
    with pytest.raises(InputError, match="span"):
        read_transactions(path)


def test_xlsx_nonempty_extra_cells_and_nontext_headers(write_input, record):
    path = write_input([record()], format="xlsx")
    book = load_workbook(path)
    book["Master"]["G2"] = "not declared"
    book.save(path)
    book.close()
    with pytest.raises(InputError, match="beyond"):
        read_transactions(path)
    path = write_input([record()], format="xlsx")
    book = load_workbook(path)
    book["Master"]["A1"] = 1
    book.save(path)
    book.close()
    with pytest.raises(InputError, match="Header"):
        read_transactions(path)


@pytest.mark.parametrize(
    "field,value,code",
    [
        ("transaction_id", 1, "ID_NOT_TEXT"),
        ("amount", 12.34, "INVALID_CELL_TYPE"),
        ("amount", True, "INVALID_CELL_TYPE"),
        ("transaction_date", date(2024, 1, 2), "INVALID_CELL_TYPE"),
        ("transaction_date", datetime(2024, 1, 2), "INVALID_CELL_TYPE"),
        ("amount", "=1+1", "FORMULA_NOT_ALLOWED"),
        ("amount", "#DIV/0!", "EXCEL_CELL_ERROR"),
    ],
)
def test_s14_nontext_business_cells_rejected_without_inference(
    write_input, record, field, value, code
):
    path = write_input([record()], format="xlsx")
    book = load_workbook(path)
    cell = book["Master"].cell(2, list(BUSINESS_COLUMNS).index(field) + 1)
    cell.value = value
    if value == "#DIV/0!":
        cell.data_type = "e"
    book.save(path)
    book.close()
    result = analyze_transactions(path)
    item = result.classified.rejected[0]
    assert code in {issue.code for issue in item.issues}
    assert item.row.raw.cells[field].kind != "text"
    assert result.quality.unquantified_count == int(field == "amount")


@pytest.mark.parametrize(
    "provenance,coverage,warning",
    [
        ({}, "absent", None),
        ({"source_file": "", "source_sheet": "", "source_row": ""}, "absent", None),
        ({"source_file": "  ", "source_row": " \t"}, "absent", None),
        (
            {"source_file": "original.csv", "source_sheet": "", "source_row": "0002"},
            "complete",
            None,
        ),
        (
            {"source_file": "original.xlsx", "source_sheet": "Sales", "source_row": "2"},
            "complete",
            None,
        ),
        ({"source_file": "original.csv"}, "partial", "INCOMPLETE_PROVENANCE"),
        ({"source_sheet": "Sales"}, "partial", "INCOMPLETE_PROVENANCE"),
        ({"source_row": "2"}, "partial", "INCOMPLETE_PROVENANCE"),
        ({"source_file": "original.csv", "source_row": "0"}, "invalid", "INVALID_PROVENANCE"),
        ({"source_file": "original.csv", "source_row": "-1"}, "invalid", "INVALID_PROVENANCE"),
        ({"source_file": "original.csv", "source_row": "2.0"}, "invalid", "INVALID_PROVENANCE"),
        ({"source_file": "original.csv", "source_row": "bad"}, "invalid", "INVALID_PROVENANCE"),
    ],
)
def test_s13_provenance_is_optional_unverified_and_independent(
    write_input, record, transport, provenance, coverage, warning
):
    headers = [*BUSINESS_COLUMNS, *provenance]
    row = record(**provenance)
    result = analyze_transactions(write_input([row], format=transport, headers=headers))
    item = result.classified.accepted[0]
    assert item.row.provenance.coverage == coverage
    assert getattr(result.quality.provenance, coverage) == 1
    assert (
        dict(item.row.raw.raw_values) == row
        if transport == "csv"
        else all(
            item.row.raw.raw_values[key] == value
            or value == ""
            and item.row.raw.raw_values[key] is None
            for key, value in row.items()
        )
    )
    assert result.quality.review_required == (warning is not None)
    assert result.quality.review_reasons == (() if warning is None else (warning,))
    assert item.row.amount_minor == 1234
    assert item.row.raw.location.import_file == f"sales.{transport}"
    assert result.upstream is None
    if coverage == "complete":
        assert item.row.provenance.source_row == 2


@pytest.mark.parametrize(
    "source_row,valid",
    [(2, True), (2.0, True), (0, False), (-1, False), (2.5, False), (True, False)],
)
def test_provenance_numeric_row_exception(write_input, record, source_row, valid):
    headers = [*BUSINESS_COLUMNS, "source_file", "source_row"]
    result = analyze_transactions(
        write_input(
            [
                record(source_file="original.csv", source_row=source_row),
            ],
            headers=headers,
            format="xlsx",
        )
    )
    assert result.quality.counts.accepted == 1
    assert result.classified.accepted[0].row.provenance.coverage == (
        "complete" if valid else "invalid"
    )
    assert result.quality.review_required != valid


def test_invalid_provenance_type_warns_not_rejects(write_input, record):
    result = analyze_transactions(
        write_input(
            [
                record(source_file=123, source_sheet=True, source_row="2"),
            ],
            headers=[*BUSINESS_COLUMNS, "source_file", "source_sheet", "source_row"],
            format="xlsx",
        )
    )
    assert result.quality.counts.accepted == 1
    warnings = result.classified.accepted[0].row.warnings
    assert [(i.field, i.code) for i in warnings] == [
        ("source_file", "INVALID_PROVENANCE"),
        ("source_sheet", "INVALID_PROVENANCE"),
    ]


def test_stale_dimensions_are_not_trusted(write_input, record):
    path = write_input([record(), record(transaction_id="two")], format="xlsx")
    with ZipFile(path) as source:
        members = {name: source.read(name) for name in source.namelist()}
    members["xl/worksheets/sheet1.xml"] = members["xl/worksheets/sheet1.xml"].replace(
        b'<dimension ref="A1:F3"', b'<dimension ref="A1:A1"'
    )
    with ZipFile(path, "w", compression=ZIP_DEFLATED) as target:
        for name, data in members.items():
            target.writestr(name, data)
    assert len(read_transactions(path).records) == 2


def test_unsupported_zip_compression_is_structural_error(write_input, record):
    path = write_input([record()], format="xlsx")
    with ZipFile(path) as source:
        members = {name: source.read(name) for name in source.namelist()}
    with ZipFile(path, "w", compression=ZIP_BZIP2) as target:
        for name, data in members.items():
            target.writestr(name, data)
    with pytest.raises(InputError, match="compression"):
        read_transactions(path)


def test_symlink_inputs_are_refused_if_available(write_input, record, tmp_path):
    source = write_input([record()])
    link = tmp_path / "linked.csv"
    try:
        link.symlink_to(source)
    except (OSError, NotImplementedError):
        pytest.skip("Creating symlinks requires OS permissions")
    with pytest.raises(InputError, match="links"):
        read_transactions(link)


def test_large_csv_text_has_only_the_documented_file_budget(write_input, record):
    identifier = "零" * 150_000
    path = write_input([record(customer_id=identifier)])
    previous = csv.field_size_limit()
    result = analyze_transactions(path)
    assert result.classified.accepted[0].row.customer_id == identifier
    assert csv.field_size_limit() == previous


def test_long_source_row_zero_padding_is_not_a_false_warning(write_input, record):
    path = write_input(
        [
            record(source_file="source.csv", source_row="0" * 4500 + "2"),
        ],
        headers=[*BUSINESS_COLUMNS, "source_file", "source_row"],
    )
    result = analyze_transactions(path)
    assert result.classified.accepted[0].row.provenance.source_row == 2
    assert not result.quality.review_required


def test_read_detects_file_changes_during_snapshot(write_input, record, monkeypatch):
    path = write_input([record()])
    real_fstat = readers.os.fstat
    calls = 0

    def changed_fstat(descriptor):
        nonlocal calls
        actual = real_fstat(descriptor)
        calls += 1
        return SimpleNamespace(
            st_mode=actual.st_mode,
            st_dev=actual.st_dev,
            st_ino=actual.st_ino,
            st_size=actual.st_size,
            st_mtime_ns=actual.st_mtime_ns + calls,
        )

    monkeypatch.setattr(readers.os, "fstat", changed_fstat)
    with pytest.raises(InputError, match="changed while reading"):
        read_transactions(path)


def test_link_guard_without_platform_symlink_privileges(write_input, record, monkeypatch):
    path = write_input([record()])
    real_is_symlink = Path.is_symlink
    monkeypatch.setattr(Path, "is_symlink", lambda self: self == path or real_is_symlink(self))
    with pytest.raises(InputError, match="links"):
        read_transactions(path)


def test_zip_without_a_workbook_is_a_public_input_error(tmp_path):
    path = tmp_path / "empty.xlsx"
    with ZipFile(path, "w"):
        pass
    with pytest.raises(InputError):
        read_transactions(path)


def test_xlsx_sparse_record_limit_is_inclusive(write_input, record):
    path = write_input([record(), record(transaction_id="two")], format="xlsx")
    assert len(read_transactions(path, limits=InputLimits(max_records=2)).records) == 2
    with pytest.raises(InputError, match="max_records"):
        read_transactions(path, limits=InputLimits(max_records=1))
