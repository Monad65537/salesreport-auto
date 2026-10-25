"""Opt-in tests use the sibling CLI only as setup, with all outputs in pytest temp."""

import hashlib
import json
import os
import subprocess
from collections import Counter
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from salesreport_auto import BundleError, analyze_transactions, generate_report, verify_workbook

pytestmark = pytest.mark.sheetmerge_integration


@pytest.fixture(scope="module")
def sibling_environment():
    sibling = Path(__file__).resolve().parents[2] / "sheetmerge"
    python = sibling / ".venv" / ("Scripts/python.exe" if os.name == "nt" else "bin/python")
    if not python.is_file() or not (sibling / "samples/input/processable").is_dir():
        pytest.skip("Sibling SheetMerge fixtures/virtual-environment Python are unavailable")
    environment = os.environ.copy()
    environment["PYTHONPATH"] = str(sibling / "src")
    environment["PYTHONDONTWRITEBYTECODE"] = "1"
    probe = subprocess.run(
        [
            str(python),
            "-B",
            "-c",
            "import sys, pandas, openpyxl; assert sys.version_info >= (3, 12)",
        ],
        env=environment,
        capture_output=True,
        text=True,
        check=False,
    )
    if probe.returncode != 0:
        pytest.skip("Sibling Python 3.12+/pandas/openpyxl environment is incompatible")
    return sibling, python, environment


def _run(sibling_environment, source, output):
    sibling, python, environment = sibling_environment
    return subprocess.run(
        [
            str(python),
            "-B",
            "-m",
            "sheetmerge",
            "--config",
            str(sibling / "samples/sources.toml"),
            "--input",
            str(sibling / source),
            "--output",
            str(output),
        ],
        env=environment,
        capture_output=True,
        text=True,
        check=False,
    )


@pytest.fixture(scope="module")
def real_bundle(sibling_environment, tmp_path_factory):
    sibling, _, _ = sibling_environment
    sources = sorted(
        path for path in (sibling / "samples/input").rglob("*") if path.suffix in (".csv", ".xlsx")
    )
    before = {
        path: (hashlib.sha256(path.read_bytes()).hexdigest(), path.stat().st_mtime_ns)
        for path in sources
    }
    destination = tmp_path_factory.mktemp("real-sheetmerge") / "bundle"
    process = _run(sibling_environment, "samples/input/processable", destination)
    assert process.returncode == 1, process.stderr
    assert json.loads(process.stdout)["counts"] == {
        "input": 33,
        "accepted": 16,
        "rejected": 14,
        "redundant": 3,
    }
    yield destination
    assert before == {
        path: (hashlib.sha256(path.read_bytes()).hexdigest(), path.stat().st_mtime_ns)
        for path in sources
    }


def _table(book, sheet_name, name):
    sheet = book[sheet_name]
    bounds = range_boundaries(sheet.tables[name].ref)
    rows = list(
        sheet.iter_rows(
            min_col=bounds[0],
            min_row=bounds[1],
            max_col=bounds[2],
            max_row=bounds[3],
            values_only=True,
        )
    )
    return [dict(zip(rows[0], row, strict=True)) for row in rows[1:]]


@pytest.mark.parametrize("master", ["master.csv", "master.xlsx"])
def test_real_sheetmerge_bundle_and_standalone_agree(real_bundle, tmp_path, master):
    selected = real_bundle / master
    report = generate_report(selected, tmp_path / "bundled.xlsx", bundle=real_bundle)
    source = report.analysis.upstream.reconciliation
    assert (
        source.counts.input,
        source.counts.accepted,
        source.counts.rejected,
        source.counts.redundant,
    ) == (33, 16, 14, 3)
    assert source.unquantified_count == 5 and source.unknown_currency_count == 0
    assert not source.monetary_complete
    assert report.analysis.quality.counts.input == report.analysis.quality.counts.accepted == 16
    assert report.analysis.quality.counts.rejected == report.analysis.quality.counts.redundant == 0
    assert report.status == "review_required"
    assert report.review_reasons == ("UPSTREAM_REJECTIONS", "UPSTREAM_UNQUANTIFIED_AMOUNTS")
    assert {
        row.currency: row.metrics.net_revenue for row in report.analysis.aggregates.summary
    } == {
        "USD": "141.45",
        "PLN": "220.35",
        "JPY": "1290",
    }
    assert {
        row.currency: row.metrics.transaction_count for row in report.analysis.aggregates.summary
    } == {"USD": 7, "PLN": 5, "JPY": 4}
    expected_reasons = {
        "CONFLICTING_TRANSACTION": 6,
        "EXACT_DUPLICATE": 3,
        "EXCESS_FRACTIONAL_PRECISION": 2,
        "KEY_HAS_INVALID_CANDIDATE": 2,
        "INVALID_DATE": 3,
        "MISSING_REQUIRED_VALUE": 3,
    }
    assert dict(report.analysis.upstream.reason_counts) == expected_reasons
    standalone = generate_report(selected, tmp_path / "standalone.xlsx")
    assert standalone.analysis.aggregates == report.analysis.aggregates
    assert standalone.analysis.upstream is None and standalone.status == "completed"
    assert standalone.to_summary()["upstream"] is None
    book = load_workbook(report.output)
    try:
        assert book.sheetnames == [
            "Summary",
            "Monthly Sales",
            "Branches",
            "Data Quality",
            "Raw Data",
        ]
        assert len(_table(book, "Raw Data", "ImportedRecords")) == 16
        assert len(_table(book, "Data Quality", "QualityDetails")) == 17
        assert Counter(
            row["status"] for row in _table(book, "Data Quality", "UpstreamExclusions")
        ) == {"rejected": 14, "redundant": 3}
        upstream = _table(book, "Data Quality", "UpstreamQuality")[0]
        assert upstream["input"] == 33 and upstream["unquantified_count"] == 5
        ledgers = {row["currency"]: row for row in _table(book, "Data Quality", "UpstreamLedgers")}
        assert [
            (
                ledgers[code]["input_known_amount"],
                ledgers[code]["accepted_known_amount"],
                ledgers[code]["rejected_known_amount"],
                ledgers[code]["redundant_known_amount"],
            )
            for code in ("USD", "PLN", "JPY")
        ] == [
            ("356.55", "141.45", "115.00", "100.10"),
            ("583.35", "220.35", "163.00", "200.00"),
            ("7389", "1290", "5099", "1000"),
        ]
        for name in ("Monthly Sales", "Branches"):
            assert len(book[name]._charts) == 3
            for chart, code in zip(book[name]._charts, ("JPY", "PLN", "USD"), strict=True):
                area = range_boundaries(chart.series[0].val.numRef.f.split("!")[1])
                assert all(
                    book[name].cell(row, 10).value == code for row in range(area[1], area[3] + 1)
                )
    finally:
        book.close()
    verify_workbook(report.output, report.plan)
    assert analyze_transactions(selected).upstream is None


def test_real_sheetmerge_structural_failure_publishes_no_bundle(sibling_environment, tmp_path):
    output = tmp_path / "must-not-exist"
    process = _run(sibling_environment, "samples/input", output)
    assert process.returncode == 2
    assert "account_id" in process.stderr and "currency_code" in process.stderr
    assert not output.exists()
    report = tmp_path / "must-not-publish.xlsx"
    with pytest.raises(BundleError):
        generate_report(output / "master.csv", report, bundle=output)
    assert not report.exists()
